from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import json
from types import SimpleNamespace

import pandas as pd
from openpyxl import load_workbook

import app_validacion as app
from extractor_metadata import MetadataEmpresa


def _archivo(nombre: str, contenido: bytes) -> BytesIO:
    archivo = BytesIO(contenido)
    archivo.name = nombre
    return archivo


def _documento(
        document_id: str, contenido: bytes, anio: int | None,
) -> dict[str, object]:
    return {
        "document_id": document_id,
        "status": "AUTHORIZED",
        "sha256": sha256(contenido).hexdigest(),
        "confirmed_closing_year": anio,
        "closing_year_evidence": "visible-document-heading",
    }


def _manifiesto(
        documentos: list[dict[str, object]], *, indent: int | None = None,
) -> BytesIO:
    contenido = json.dumps(
        {
            "schema_version": "1.0",
            "pilot": {
                "pilot_id": "pilot-opaque-001",
                "client_id": "client-opaque-001",
                "max_authorized_documents": 10,
            },
            "documents": documentos,
        },
        indent=indent,
        sort_keys=True,
    ).encode("utf-8")
    return _archivo("manifest.json", contenido)


def _estado() -> dict:
    return {
        "document_pages": {},
        "document_closing_years": {},
        "document_metadata_confirmations": {},
        "resultados": {},
        "metadata_files": {},
        "document_intel": {},
        "extraction_pending": {},
        "extraction_resolved": {},
        "extraction_certifications": {},
        "processed_at": {},
        "quality_controls": {},
        "delivery_readiness": {},
        "historial_decisiones": [],
        "extraction_revisions": {},
    }


def _errores_silenciosos(monkeypatch) -> list[str]:
    mensajes: list[str] = []
    monkeypatch.setattr(
        app.st, "error", lambda mensaje, *args, **kwargs: mensajes.append(mensaje),
    )
    monkeypatch.setattr(app.st, "caption", lambda *args, **kwargs: None)
    return mensajes


def test_manifiesto_piloto_asocia_bytes_en_orden_invertido_y_fija_anios(
        monkeypatch,
):
    monkeypatch.setenv("PILOT_MODE", "1")
    mensajes = _errores_silenciosos(monkeypatch)
    state = _estado()
    monkeypatch.setattr(app.st, "session_state", state)
    primero = _archivo("nombre-visible-primero.pdf", b"contenido primero")
    segundo = _archivo("nombre-visible-segundo.pdf", b"contenido segundo")
    manifiesto = _manifiesto([
        _documento("doc-opaque-primero", primero.getvalue(), 2021),
        _documento("doc-opaque-segundo", segundo.getvalue(), None),
    ])

    assert app._autorizar_lote_piloto([segundo, primero], manifiesto)
    assert not mensajes
    assert state["pilot_manifest_authorizations"][segundo.name] == {
        "document_id": "doc-opaque-segundo",
        "confirmed_closing_year": None,
        "source_sha256": sha256(segundo.getvalue()).hexdigest(),
    }
    assert state["pilot_manifest_authorizations"][primero.name] == {
        "document_id": "doc-opaque-primero",
        "confirmed_closing_year": 2021,
        "source_sha256": sha256(primero.getvalue()).hexdigest(),
    }

    assert app._confirmar_anios_cierre_documentos([segundo, primero])
    assert state["document_closing_years"][segundo.name]["anio_cierre"] is None
    assert state["document_closing_years"][segundo.name]["confirmado"] is True
    assert state["document_closing_years"][primero.name]["anio_cierre"] == 2021
    assert state["document_closing_years"][primero.name]["origen"] == "pilot_manifest"


def test_manifiesto_con_documento_extra_bloquea_y_limpia_estado_derivado(
        monkeypatch,
):
    monkeypatch.setenv("PILOT_MODE", "1")
    mensajes = _errores_silenciosos(monkeypatch)
    state = _estado()
    conocido = _archivo("nombre-visible.pdf", b"contenido autorizado")
    extra = _archivo("otro-visible.pdf", b"contenido no autorizado")
    source_sha = sha256(conocido.getvalue()).hexdigest()
    state["pilot_manifest_context"] = {
        "manifest_sha256": "a" * 64,
        "authorizations": {
            conocido.name: {
                "document_id": "doc-opaque-conocido",
                "confirmed_closing_year": 2024,
                "source_sha256": source_sha,
            },
        },
    }
    state["pilot_manifest_authorizations"] = dict(
        state["pilot_manifest_context"]["authorizations"],
    )
    state["resultados"][conocido.name] = pd.DataFrame([
        {"codigo_clasificado": "AC.01"},
    ])
    state["document_closing_years"][conocido.name] = {
        **app._contexto_confirmacion_anio(conocido),
        "anio_cierre": 2024,
        "confirmado": True,
    }
    monkeypatch.setattr(app.st, "session_state", state)
    manifiesto = _manifiesto([
        _documento("doc-opaque-conocido", conocido.getvalue(), 2024),
    ])

    assert not app._autorizar_lote_piloto([conocido, extra], manifiesto)
    assert conocido.name not in state["resultados"]
    assert conocido.name not in state["document_closing_years"]
    assert state["pilot_manifest_authorizations"] == {}
    assert "pilot_manifest_context" not in state
    assert mensajes == [
        "No se procesó el lote. Códigos: unexpected_uploaded_content",
    ]
    assert conocido.name not in mensajes[0]
    assert extra.name not in mensajes[0]
    assert source_sha not in mensajes[0]


def test_manifiesto_con_documento_faltante_indica_solo_las_cantidades(
        monkeypatch,
):
    monkeypatch.setenv("PILOT_MODE", "1")
    mensajes = _errores_silenciosos(monkeypatch)
    state = _estado()
    monkeypatch.setattr(app.st, "session_state", state)
    primero = _archivo("primero-visible.pdf", b"primer contenido")
    segundo = _archivo("segundo-visible.pdf", b"segundo contenido")
    manifiesto = _manifiesto([
        _documento("doc-opaque-primero", primero.getvalue(), 2023),
        _documento("doc-opaque-segundo", segundo.getvalue(), 2024),
    ])

    assert not app._autorizar_lote_piloto([primero], manifiesto)
    assert mensajes == [
        "No se procesó el lote. Códigos: missing_uploaded_document. "
        "El manifiesto autoriza 2 documento(s) y se cargaron 1. "
        "Cargue todos los documentos autorizados juntos.",
    ]
    assert "primero-visible.pdf" not in mensajes[0]
    assert "doc-opaque-segundo" not in mensajes[0]


def test_cambio_de_manifiesto_invalida_solo_autorizacion_efectiva_distinta(
        monkeypatch,
):
    monkeypatch.setenv("PILOT_MODE", "1")
    _errores_silenciosos(monkeypatch)
    state = _estado()
    monkeypatch.setattr(app.st, "session_state", state)
    primero = _archivo("primero-visible.pdf", b"primer contenido")
    segundo = _archivo("segundo-visible.pdf", b"segundo contenido")
    documentos = [
        _documento("doc-opaque-primero", primero.getvalue(), 2022),
        _documento("doc-opaque-segundo", segundo.getvalue(), 2024),
    ]

    assert app._autorizar_lote_piloto([primero, segundo], _manifiesto(documentos))
    state["resultados"] = {
        primero.name: pd.DataFrame([{"codigo_clasificado": "AC.01"}]),
        segundo.name: pd.DataFrame([{"codigo_clasificado": "PAT.01"}]),
    }
    state["historial_decisiones"] = [
        {"Archivo": primero.name, "Cuenta": "Caja"},
        {"Archivo": segundo.name, "Cuenta": "Capital"},
    ]

    # El hash del JSON cambia por formato, pero la autorización efectiva no.
    assert app._autorizar_lote_piloto(
        [primero, segundo], _manifiesto(documentos, indent=2),
    )
    assert set(state["resultados"]) == {primero.name, segundo.name}

    documentos[0] = _documento(
        "doc-opaque-primero", primero.getvalue(), 2021,
    )
    assert app._autorizar_lote_piloto([primero, segundo], _manifiesto(documentos))
    assert primero.name not in state["resultados"]
    assert segundo.name in state["resultados"]
    assert state["historial_decisiones"] == [
        {"Archivo": segundo.name, "Cuenta": "Capital"},
    ]


def test_payload_piloto_incluye_id_y_hash_y_fuera_de_piloto_no_cambia(
        monkeypatch,
):
    state = _estado()
    source_sha = sha256(b"contenido autorizado").hexdigest()
    state["pilot_manifest_authorizations"] = {
        "nombre-visible.pdf": {
            "document_id": "doc-opaque-991",
            "confirmed_closing_year": 2024,
            "source_sha256": source_sha,
        },
    }
    monkeypatch.setattr(app.st, "session_state", state)
    estado = {
        "estado": "REVISION", "etiqueta": "Revisión",
        "exportacion_final_permitida": False,
        "exportacion_revision_permitida": True,
        "razones": [],
    }

    monkeypatch.setenv("PILOT_MODE", "1")
    piloto = app._payload_entrega_json(
        archivo_nombre="nombre-visible.pdf", meta=None, estado_entrega=estado,
        emision={"definitivo": False, "motivos": []},
        presentacion=pd.DataFrame(), detalle=pd.DataFrame(),
    )
    assert piloto["archivo_origen"] == "doc-opaque-991"
    assert piloto["documento_piloto"] == {
        "document_id": "doc-opaque-991", "source_sha256": source_sha,
    }
    assert piloto["ejecucion"]["modo_piloto"] is True
    assert piloto["ejecucion"]["rama"]
    assert piloto["ejecucion"]["commit"]

    monkeypatch.delenv("PILOT_MODE", raising=False)
    normal = app._payload_entrega_json(
        archivo_nombre="nombre-visible.pdf", meta=None, estado_entrega=estado,
        emision={"definitivo": False, "motivos": []},
        presentacion=pd.DataFrame(), detalle=pd.DataFrame(),
    )
    assert normal["archivo_origen"] == "nombre-visible.pdf"
    assert "documento_piloto" not in normal


def test_excel_piloto_usa_referencia_opaca_y_no_filtra_nombre_del_historial(
        monkeypatch,
):
    from streamlit.testing.v1 import AppTest

    monkeypatch.setenv("PILOT_MODE", "1")
    monkeypatch.setenv("QUALITY_CONTROL_ENFORCE_EXPORT", "false")
    source_name = "nombre-visible.pdf"
    source_sha = sha256(b"contenido autorizado").hexdigest()
    filas = [
        {
            "nombre_original": "Caja", "codigo_original": "1101-01",
            "monto": 100, "origen_columna": "activo",
            "origen_columna_efectiva": "activo", "codigo_clasificado": "AC.01",
            "es_total": False, "requiere_revision": False,
            "metodo": "validacion_humana", "confianza": 1.0,
            "nombre_revision_usuario": "",
        },
        {
            "nombre_original": "Capital", "codigo_original": "2401-01",
            "monto": 100, "origen_columna": "pasivo",
            "origen_columna_efectiva": "patrimonio", "codigo_clasificado": "PAT.01",
            "es_total": False, "requiere_revision": False,
            "metodo": "validacion_humana", "confianza": 1.0,
            "nombre_revision_usuario": "",
        },
    ]

    def run(rows, nombre, digest):
        import pandas as pandas_local
        import streamlit as st
        import app_validacion as app_local
        from types import SimpleNamespace as SimpleNamespaceLocal
        from extractor_metadata import MetadataEmpresa as MetadataEmpresaLocal
        from reporting_integrity import catalogo_local

        st.session_state.metadata_files = {
            nombre: MetadataEmpresaLocal(
                rut="76.000.001-1", razon_social="Razón visible", moneda="$",
            ),
        }
        st.session_state.extraction_certifications = {
            nombre: SimpleNamespaceLocal(
                estado="certificada", metodo="coordinates_8_amounts",
                orientacion_certificada=True, razones=[],
            ),
        }
        st.session_state.pilot_manifest_authorizations = {
            nombre: {
                "document_id": "doc-opaque-991",
                "confirmed_closing_year": 2024,
                "source_sha256": digest,
            },
        }
        st.session_state.historial_decisiones = [{
            "Archivo": nombre, "Fila": "0", "Cuenta": "Caja",
            "Clasificación anterior": "", "Clasificación nueva": "AC.01",
            "Método anterior": "manual", "Motivo": "prueba", "Fecha": "2026-10-15",
        }]
        previo = app_local.st.download_button

        def capturar(label, **kwargs):
            exportaciones = list(st.session_state.get("exportaciones", []))
            exportaciones.append({"label": label, **kwargs})
            st.session_state.exportaciones = exportaciones

        app_local.st.download_button = capturar
        try:
            app_local._tab_balance(
                pandas_local.DataFrame(rows), catalogo_local(), nombre,
            )
        finally:
            app_local.st.download_button = previo

    at = AppTest.from_function(
        run, args=(filas, source_name, source_sha),
    ).run(timeout=20)
    assert not at.exception
    exportaciones = at.session_state.exportaciones
    assert len(exportaciones) == 2
    for descarga in exportaciones:
        assert descarga["file_name"].startswith(
            f"PILOTO-BALANCE-doc-opaque-991-{source_sha}",
        )
        assert source_name not in descarga["file_name"]
        assert "760000011" not in descarga["file_name"]

    descarga_json = next(
        descarga for descarga in exportaciones
        if descarga["mime"] == "application/json"
    )
    payload = json.loads(descarga_json["data"].decode("utf-8"))
    assert payload["documento_piloto"] == {
        "document_id": "doc-opaque-991", "source_sha256": source_sha,
    }
    assert payload["archivo_origen"] == "doc-opaque-991"

    descarga_excel = next(
        descarga for descarga in exportaciones
        if descarga["mime"].startswith("application/vnd.openxmlformats")
    )
    libro = load_workbook(BytesIO(descarga_excel["data"]), data_only=True)
    controles = {
        fila[0]: fila[1]
        for fila in libro["Control de emisión"].iter_rows(min_row=2, values_only=True)
        if fila[0]
    }
    assert controles["ID opaco de documento"] == "doc-opaque-991"
    assert controles["SHA-256 de origen"] == source_sha
    assert controles["Rama de ejecución"]
    assert controles["Commit de ejecución"]
    decisiones = list(libro["Decisiones de esta sesión"].values)
    encabezado = decisiones[0]
    fila_decision = decisiones[1]
    assert fila_decision[encabezado.index("Archivo")] == "doc-opaque-991"
    assert source_name not in " ".join(
        "" if valor is None else str(valor)
        for fila in decisiones for valor in fila
    )
