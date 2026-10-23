"""Contrato de salida que no depende del recálculo de un libro Excel."""

from io import BytesIO
import json

from openpyxl import load_workbook
import pandas as pd
from streamlit.testing.v1 import AppTest

import app_validacion as app
from extractor_metadata import MetadataEmpresa
from test_report_integrity import afuminsal


def _app_exportacion(rows, meta, certificacion=None):
    import app_validacion as app
    import pandas as pandas_local
    import streamlit as st
    from reporting_integrity import catalogo_local as catalogo_local_func
    from types import SimpleNamespace as SimpleNamespaceLocal

    nombre = "caso.pdf"
    st.session_state.metadata_files = {nombre: meta}
    certificacion = certificacion or {
        "estado": "certificada",
        "metodo": "coordinates_8_amounts",
        "orientacion_certificada": True,
        "razones": [],
    }
    st.session_state.extraction_certifications = {
        nombre: SimpleNamespaceLocal(**certificacion),
    }
    previo = app.st.download_button

    def capturar(label, **kwargs):
        exportaciones = list(st.session_state.get("exportaciones", []))
        exportaciones.append({"label": label, **kwargs})
        st.session_state.exportaciones = exportaciones

    app.st.download_button = capturar
    try:
        app._tab_balance(
            pandas_local.DataFrame(rows), catalogo_local_func(), nombre,
        )
    finally:
        app.st.download_button = previo


def _controles(libro):
    return {
        fila[0]: fila[1]
        for fila in libro["Control de emisión"].iter_rows(
            min_row=2, values_only=True,
        )
        if fila[0]
    }


def _fila(
        nombre, monto, origen, codigo, *, estado_financiero=None,
) -> dict:
    fila = {
        "nombre_original": nombre,
        "codigo_original": nombre,
        "nombre_revision_usuario": "",
        "monto": monto,
        "origen_columna": origen,
        "origen_columna_efectiva": origen,
        "codigo_clasificado": codigo,
        "es_total": False,
        "requiere_revision": False,
        "metodo": "code",
        "confianza": 1.0,
    }
    if estado_financiero is not None:
        fila["estado_financiero"] = estado_financiero
    return fila


def test_paquete_exportado_tiene_controles_estaticos_y_metadata_explicitamente_ausente(
        monkeypatch,
):
    monkeypatch.setenv("QUALITY_CONTROL_ENFORCE_EXPORT", "false")
    monkeypatch.delenv("PILOT_MODE", raising=False)
    at = AppTest.from_function(
        _app_exportacion,
        args=(afuminsal(False), MetadataEmpresa(moneda="$")),
    ).run(timeout=20)

    assert not at.exception
    exportaciones = at.session_state.exportaciones
    descarga_excel = next(
        descarga for descarga in exportaciones
        if descarga["mime"].startswith(
            "application/vnd.openxmlformats-officedocument",
        )
    )
    descarga_json = next(
        descarga for descarga in exportaciones
        if descarga["mime"] == "application/json"
    )

    # Un consumidor que no recalcula fórmulas debe obtener los controles
    # críticos directamente desde esta hoja.
    libro_valores = load_workbook(
        BytesIO(descarga_excel["data"]), data_only=True,
    )
    controles = _controles(libro_valores)
    assert controles["Contrato de controles"] == (
        "Valores materializados por el proceso; no requieren recálculo del libro."
    )
    assert controles[
        "Requiere recálculo del libro para estos controles"
    ] is False
    assert controles["Total activos (materializado)"] == 29_943_255
    assert controles["Total pasivos (materializado)"] == 3_177_262
    assert controles["Total patrimonio (materializado)"] == 26_765_993
    assert controles["Pasivo más patrimonio (materializado)"] == 29_943_255
    assert controles[
        "Diferencia Activo menos Pasivo y Patrimonio (materializada)"
    ] == 0
    assert controles["Empresa (metadata documental)"] == "No informado"
    assert controles["RUT (metadata documental)"] == "No informado"
    assert controles["Analista (ID autenticado)"] == "No informado"
    assert libro_valores["Balance Normalizado"]["B1"].value == "No informado"
    assert libro_valores["Resumen"]["B7"].value == "No informado"

    # La misma hoja no usa fórmulas para los controles que el contrato declara
    # materializados, aunque las hojas de presentación sigan siendo auditables.
    libro_formulas = load_workbook(
        BytesIO(descarga_excel["data"]), data_only=False,
    )
    controles_formula = _controles(libro_formulas)
    for etiqueta in (
        "Total activos (materializado)",
        "Total pasivos (materializado)",
        "Total patrimonio (materializado)",
        "Pasivo más patrimonio (materializado)",
        "Diferencia Activo menos Pasivo y Patrimonio (materializada)",
        "Estado de cuadratura (materializado)",
    ):
        valor = controles_formula[etiqueta]
        assert not (isinstance(valor, str) and valor.startswith("="))

    payload = json.loads(descarga_json["data"].decode("utf-8"))
    assert payload["identificacion_caso"]["empresa"] == {
        "valor": None, "estado": "no_informado",
    }
    assert payload["identificacion_caso"]["rut"] == {
        "valor": None, "estado": "no_informado",
    }
    assert payload["identificacion_caso"]["analista_id"] == {
        "valor": None, "estado": "no_informado",
    }
    assert payload["controles_materializados"] == {
        "materializados": True,
        "requiere_recalculo_del_libro": False,
        "fuente_cuadratura": "certificacion_homologada_canonica",
        "total_activos": 29_943_255.0,
        "total_pasivos": 3_177_262.0,
        "total_patrimonio": 26_765_993.0,
        "patrimonio_sin_resultado": 16_855_048.0,
        "resultado_periodo_incorporado": 9_910_945.0,
        "pasivo_mas_patrimonio": 29_943_255.0,
        "diferencia_activo_menos_pasivo_patrimonio": 0.0,
        "tolerancia_cuadratura": 1.0,
        "estado_cuadratura": "CUADRA_DENTRO_DE_TOLERANCIA",
        "resultado_derivado_incorporado": True,
        "resultado_final_fuente_incorporado": False,
        "resultado_final_fuente": None,
        "resultado_segun_columnas_originales": 9_910_945.0,
        "resultado_segun_categorias_homologadas": 9_910_945.0,
        "diferencia_resultado": 0.0,
    }
    assert controles["Fuente de cuadratura materializada"] == (
        "certificacion_homologada_canonica"
    )
    assert controles["Umbral de cobertura comercial"] == 0.6
    assert controles[
        "Numerador comercial: automáticas estrictas más sugeridas automatizadas"
    ] == 0
    assert controles["Universo elegible comercial"] == 9
    assert controles["Decisiones humanas fuera del indicador comercial"] == 9
    assert controles["Pendientes comerciales"] == 0
    assert payload["estado_entrega"]["aceptacion_comercial"]["cobertura"][
        "cuentas"
    ]["automaticas_o_sugeridas"] == 0


def test_paquete_piloto_no_serializa_aceptacion_comercial_final(monkeypatch):
    monkeypatch.setenv("PILOT_MODE", "1")
    estado = {
        "estado": "LISTO_PARA_INTEGRAR",
        "etiqueta": "Listo para integrar",
        "exportacion_final_permitida": True,
        "exportacion_revision_permitida": False,
        "razones": [],
        "aceptacion_comercial": {
            "estado": "CUMPLE_CRITERIO_COMERCIAL_FINAL",
            "apto_para_revision_comercial": True,
            "apto_para_entrega_revisada": True,
            "motivos": [],
            "siguiente_accion": [],
            "cobertura": {"cuentas": {}},
        },
    }

    payload = app._payload_entrega_json(
        archivo_nombre="balance.pdf",
        meta=None,
        estado_entrega=estado,
        emision={"definitivo": True, "motivos": []},
        presentacion=pd.DataFrame(),
        detalle=pd.DataFrame(),
    )

    comercial = payload["estado_entrega"]["aceptacion_comercial"]
    assert comercial["estado"] == "PILOTO_PROVISIONAL_NO_INTEGRAR"
    assert comercial["estado_tecnico_comercial"] == (
        "CUMPLE_CRITERIO_COMERCIAL_FINAL"
    )
    assert comercial["apto_para_entrega_revisada_tecnica"] is True
    assert comercial["apto_para_entrega_revisada"] is False
    assert comercial["integracion_autorizada"] is False
    assert payload["emision"]["definitiva"] is False


def test_controles_exportados_usan_cierre_final_fuente_certificado(monkeypatch):
    monkeypatch.setenv("QUALITY_CONTROL_ENFORCE_EXPORT", "false")
    monkeypatch.delenv("PILOT_MODE", raising=False)
    filas = [
        _fila("Caja", 150, "activo", "AC.01"),
        _fila("Proveedor", 10, "pasivo", "PC.01"),
        _fila("Capital", 20, "pasivo", "PAT.01"),
        _fila("Saldo patrimonial previo", 20, "pasivo", "PAT.04"),
        _fila("Ventas", 100, "ganancia", "ER.01"),
    ]
    certificacion = {
        "estado": "certificada",
        "metodo": "ocr_classified_two_columns",
        "orientacion_certificada": True,
        "razones": [],
        "totales_finales_validos": True,
        "ecuacion_contable_fuente_valida": True,
        "resultado_ejercicio": 100,
        "tipo_resultado": "utilidad",
    }
    at = AppTest.from_function(
        _app_exportacion,
        args=(filas, MetadataEmpresa(moneda="$"), certificacion),
    ).run(timeout=20)

    assert not at.exception
    descarga_excel = next(
        descarga for descarga in at.session_state.exportaciones
        if descarga["mime"].startswith(
            "application/vnd.openxmlformats-officedocument",
        )
    )
    descarga_json = next(
        descarga for descarga in at.session_state.exportaciones
        if descarga["mime"] == "application/json"
    )
    controles = _controles(load_workbook(
        BytesIO(descarga_excel["data"]), data_only=True,
    ))
    payload = json.loads(descarga_json["data"].decode("utf-8"))

    # El diagnóstico visual de esta estructura ve sólo 50 de patrimonio.
    # El contrato debe seguir la certificación canónica: 40 patrimoniales más
    # el cierre fuente certificado de 100.
    assert controles["Total patrimonio (materializado)"] == 140
    assert controles[
        "Patrimonio sin resultado del período (materializado)"
    ] == 40
    assert controles[
        "Resultado del período incorporado (materializado)"
    ] == 100
    assert controles["Pasivo más patrimonio (materializado)"] == 150
    assert controles[
        "Diferencia Activo menos Pasivo y Patrimonio (materializada)"
    ] == 0
    assert controles["Cierre final fuente incorporado en cuadratura"] is True
    assert controles["Cierre final certificado de fuente"] == 100
    assert payload["controles_materializados"][
        "resultado_final_fuente_incorporado"
    ] is True
    assert payload["controles_materializados"][
        "resultado_final_fuente"
    ] == 100.0


def test_controles_exportados_no_duplican_resultado_de_estado_separado(monkeypatch):
    monkeypatch.setenv("QUALITY_CONTROL_ENFORCE_EXPORT", "false")
    monkeypatch.delenv("PILOT_MODE", raising=False)
    filas = [
        _fila("Caja", 100, "activo", "AC.01", estado_financiero="balance"),
        _fila(
            "Proveedor", 40, "pasivo", "PC.01", estado_financiero="balance",
        ),
        _fila("Capital", 60, "pasivo", "PAT.01", estado_financiero="balance"),
        _fila("Ventas", 30, "ganancia", "ER.01", estado_financiero="resultados"),
        _fila("Gastos", 10, "perdida", "ER.04", estado_financiero="resultados"),
    ]
    certificacion = {
        "estado": "certificada",
        "metodo": "ocr_classified_two_columns",
        "orientacion_certificada": True,
        "razones": [],
        "totales_finales_validos": True,
    }
    at = AppTest.from_function(
        _app_exportacion,
        args=(filas, MetadataEmpresa(moneda="$"), certificacion),
    ).run(timeout=20)

    assert not at.exception
    descarga_json = next(
        descarga for descarga in at.session_state.exportaciones
        if descarga["mime"] == "application/json"
    )
    payload = json.loads(descarga_json["data"].decode("utf-8"))
    controles = payload["controles_materializados"]
    assert controles["patrimonio_sin_resultado"] == 60.0
    assert controles["resultado_periodo_incorporado"] == 0.0
    assert controles["total_patrimonio"] == 60.0
    assert controles["pasivo_mas_patrimonio"] == 100.0
    assert controles["resultado_derivado_incorporado"] is False


def test_resumen_comercial_cuenta_sugerencias_fuera_del_numerador_como_pendientes():
    resumen = app._resumen_comercial_para_control({
        "aceptacion_comercial": {
            "estado": "NO_APTO_POR_COBERTURA_COMERCIAL",
            "umbral_cobertura": 0.6,
            "cobertura": {
                "cuentas": {
                    "elegibles": 3,
                    "automaticas_o_sugeridas": 0,
                    "sugerencias_metodo_no_admitido": 1,
                    "sugerencias_sin_evidencia_trazable": 1,
                    "sugerencias_confianza_insuficiente": 1,
                },
            },
        },
    })

    # El mismo conjunto alimenta el Control de emisión y la métrica de UI.
    assert resumen["pendientes"] == 3
    assert resumen["automaticas_o_sugeridas"] == 0
    assert resumen["elegibles"] == 3
