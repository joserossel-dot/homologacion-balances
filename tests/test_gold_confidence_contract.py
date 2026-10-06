"""Pruebas del contrato explícito de confianza, gobernanza y pisos de regresión en Gold Schema 2."""

from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path
import pandas as pd
import pytest

from scripts.certify_local_corpus import (
    evaluate_gold_rows,
    evaluate_expectations,
    load_gold_rows,
    _is_valid_confidence,
)
from scripts.gold_confidence_contract import (
    is_valid_sha256,
    validate_sha256_binding,
)
from scripts.migrate_gold_workbook import (
    migrate_gold_workbook,
    build_gold_review_package,
    _indexed,
)
from pipeline.homologation_pipeline import HomologationPipeline


def _base_row(**overrides) -> dict:
    row = {
        "_gold_schema_version": 2,
        "line": 1,
        "account_code": "",
        "name": "Efectivo y equivalentes al efectivo",
        "accounting_hierarchy": "AC",
        "origin": "activo",
        "amount": 1000.0,
        "period_amounts": {"2022": 1000.0},
        "column_amounts": {},
        "standard_code": "AC.01",
        "classification_method": "audited_statement_label",
        "requires_review": False,
        "is_total": False,
        "confidence": 1.0,
        "confidence_min": None,
        "derived_columns": [],
    }
    row.update(overrides)
    return row


def test_1_confianza_observada_igual_al_umbral_aprueba():
    """1. Confianza observada igual al piso de regresión por fila, aprueba sin diferencia."""
    actual = [_base_row(confidence=0.75)]
    expected = [_base_row(confidence=0.75, confidence_min=0.75)]
    results = evaluate_gold_rows(actual, expected)
    mismatch = next(c for c in results if c["name"] == "gold_mismatched_rows")
    assert mismatch["passed"] is True
    assert mismatch["actual"] == []


def test_2_confianza_observada_superior_al_umbral_aprueba():
    """2. Confianza observada superior al piso de regresión por fila, aprueba sin diferencia."""
    actual = [_base_row(confidence=0.90)]
    expected = [_base_row(confidence=0.75, confidence_min=0.75)]
    results = evaluate_gold_rows(actual, expected)
    mismatch = next(c for c in results if c["name"] == "gold_mismatched_rows")
    assert mismatch["passed"] is True
    assert mismatch["actual"] == []


def test_3_confianza_inferior_al_umbral_bloquea():
    """3. Confianza observada inferior al piso de regresión por fila, bloquea con diferencia."""
    actual = [_base_row(confidence=0.70)]
    expected = [_base_row(confidence=0.75, confidence_min=0.75)]
    results = evaluate_gold_rows(actual, expected)
    mismatch = next(c for c in results if c["name"] == "gold_mismatched_rows")
    assert mismatch["passed"] is False
    assert len(mismatch["actual"]) == 1
    diffs = mismatch["actual"][0]["differences"]
    assert "confidence" in diffs
    assert diffs["confidence"] == [0.70, 0.75]


def test_4_validacion_humana_ausente_bloquea(tmp_path: Path):
    """4. Validación humana ausente en Estado_revision, bloquea con ValueError."""
    wb_path = tmp_path / "gold_sin_estado.xlsx"
    summary = pd.DataFrame([{
        "Gold_schema_version": 2,
        "Gold_confidence_contract_version": 1,
    }])
    cuentas = pd.DataFrame([{
        "Fila": 1,
        "Cuenta_original": "Caja",
        "Origen": "activo",
        "Monto": 100,
        "Estado_revision": None,  # Ausente
        "Confianza_extraccion": 0.75,
        "Confianza_minima": 0.75,
    }])
    with pd.ExcelWriter(wb_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resumen", index=False)
        cuentas.to_excel(writer, sheet_name="Cuentas", index=False)

    with pytest.raises(ValueError, match="sin Estado_revision=APROBADO o EXCLUIR"):
        load_gold_rows(wb_path)


def test_5_estado_corregir_bloquea_certificacion(tmp_path: Path):
    """5. Estado_revision=CORREGIR bloquea con ValueError claro."""
    wb_path = tmp_path / "gold_corregir.xlsx"
    summary = pd.DataFrame([{
        "Gold_schema_version": 2,
        "Gold_confidence_contract_version": 1,
    }])
    cuentas = pd.DataFrame([{
        "Fila": 1,
        "Cuenta_original": "Caja",
        "Origen": "activo",
        "Monto": 100,
        "Estado_revision": "CORREGIR",
        "Confianza_extraccion": 0.75,
        "Confianza_minima": 0.75,
    }])
    with pd.ExcelWriter(wb_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resumen", index=False)
        cuentas.to_excel(writer, sheet_name="Cuentas", index=False)

    with pytest.raises(ValueError, match="sin Estado_revision=APROBADO o EXCLUIR"):
        load_gold_rows(wb_path)


def test_5b_estado_excluir_se_omite_y_no_genera_cuenta_aprobada(tmp_path: Path):
    """5b. Estado_revision=EXCLUIR se omite de expected rows y no genera cuenta esperada aprobada."""
    wb_path = tmp_path / "gold_excluir.xlsx"
    summary = pd.DataFrame([{
        "Gold_schema_version": 2,
        "Gold_confidence_contract_version": 1,
    }])
    cuentas = pd.DataFrame([
        {
            "Fila": 1,
            "Cuenta_original": "Caja",
            "Origen": "activo",
            "Monto": 100,
            "Estado_revision": "APROBADO",
            "Confianza_extraccion": 1.0,
            "Confianza_minima": 1.0,
        },
        {
            "Fila": 2,
            "Cuenta_original": "Ruido OCR",
            "Origen": "activo",
            "Monto": 0,
            "Estado_revision": "EXCLUIR",
            "Confianza_extraccion": 0.5,
            "Confianza_minima": 0.5,
        },
    ])
    with pd.ExcelWriter(wb_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resumen", index=False)
        cuentas.to_excel(writer, sheet_name="Cuentas", index=False)

    rows = load_gold_rows(wb_path)
    assert len(rows) == 1
    assert rows[0]["name"] == "Caja"


def test_6_campos_contables_distintos_bloquea():
    """6. Campos contables distintos (ej. monto, código), bloquea independientemente de confianza."""
    actual = [_base_row(amount=2000.0, standard_code="AC.02", confidence=0.75)]
    expected = [_base_row(amount=1000.0, standard_code="AC.01", confidence=0.75, confidence_min=0.75)]
    results = evaluate_gold_rows(actual, expected)
    mismatch = next(c for c in results if c["name"] == "gold_mismatched_rows")
    assert mismatch["passed"] is False
    diffs = mismatch["actual"][0]["differences"]
    assert "amount" in diffs
    assert "standard_code" in diffs


def test_7_gold_legado_sin_marcador_mantiene_igualdad_estricta():
    """7. Gold legado sin marcador de contrato (confidence_min=None) mantiene igualdad estricta."""
    actual = [_base_row(confidence=0.75)]
    expected = [_base_row(confidence=1.0, confidence_min=None)]
    results = evaluate_gold_rows(actual, expected)
    mismatch = next(c for c in results if c["name"] == "gold_mismatched_rows")
    assert mismatch["passed"] is False
    assert mismatch["actual"][0]["differences"]["confidence"] == [0.75, 1.0]


def test_8_migracion_conserva_todos_los_atributos_contables_y_metadatos(tmp_path: Path):
    """8. La migración explícita preserva todos los atributos contables y metadatos protegidos."""
    legacy_path = tmp_path / "legacy.xlsx"
    cand_path = tmp_path / "candidate.xlsx"
    out_path = tmp_path / "out.xlsx"

    row_data = {
        "Fila": 1,
        "Codigo_original": "100.01",
        "Cuenta_original": "Caja chica sucursal",
        "Jerarquia_contable": "AC",
        "Origen": "activo",
        "Monto": -500.5,
        "Monto_actual": -500.5,
        "Monto_anterior": -400.0,
        "Montos_periodos_JSON": '{"2021": -400.0, "2022": -500.5}',
        "Montos_columnas_JSON": '{"CLP": -500.5}',
        "Codigo_homologado": "AC.01",
        "Metodo_clasificacion": "audited_statement_label",
        "Requiere_revision": False,
        "Es_control_total": False,
        "Confianza_extraccion": 0.75,
        "Confianza_minima": 0.75,
        "Columnas_derivadas_JSON": '["periodo_actual"]',
        "Estado_revision": "APROBADO",
        "Observacion_analista": "Revisado por auditor",
    }
    valid_hash = "2c51270ea91ba1e805f99a745c217356312bf8b17d2582f40df5d6e337f30711"
    with pd.ExcelWriter(legacy_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
            "SHA256": valid_hash,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([row_data]).to_excel(writer, sheet_name="Cuentas", index=False)

    cand_row = dict(row_data)
    cand_row["Estado_revision"] = "PENDIENTE"
    cand_row["Observacion_analista"] = ""
    with pd.ExcelWriter(cand_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
            "SHA256": valid_hash,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([cand_row]).to_excel(writer, sheet_name="Cuentas", index=False)

    migrate_gold_workbook(legacy_path, cand_path, out_path)

    df_out = pd.read_excel(out_path, sheet_name="Cuentas")
    sum_out = pd.read_excel(out_path, sheet_name="Resumen")
    assert sum_out["SHA256"].iloc[0] == valid_hash

    # Verificación exhaustiva de cada campo exigido
    row = df_out.iloc[0]
    assert float(row["Monto"]) == -500.5
    assert float(row["Monto"]) < 0  # Signo negativo preservado
    assert float(row["Monto_actual"]) == -500.5
    assert float(row["Monto_anterior"]) == -400.0
    assert str(row["Codigo_original"]) == "100.01"
    assert str(row["Codigo_homologado"]) == "AC.01"
    assert str(row["Origen"]) == "activo"
    assert json.loads(row["Montos_periodos_JSON"]) == {"2021": -400.0, "2022": -500.5}
    assert json.loads(row["Montos_columnas_JSON"]) == {"CLP": -500.5}
    assert str(row["Jerarquia_contable"]) == "AC"
    assert bool(row["Es_control_total"]) is False
    assert bool(row["Requiere_revision"]) is False
    assert str(row["Metodo_clasificacion"]) == "audited_statement_label"
    assert float(row["Confianza_extraccion"]) == 0.75
    assert float(row["Confianza_minima"]) == 0.75
    assert json.loads(row["Columnas_derivadas_JSON"]) == ["periodo_actual"]
    assert str(row["Estado_revision"]) == "APROBADO"
    assert int(sum_out.iloc[0]["Gold_schema_version"]) == 2
    assert int(sum_out.iloc[0]["Gold_confidence_contract_version"]) == 1


def test_9_pnc05_queda_vinculado_a_politica_global():
    """9. La política global aprobada del 2026-09-07 vincula Otros pasivos financieros no corrientes a PNC.05."""
    pipeline = HomologationPipeline()
    result = pipeline._classify_audited_statement_label(
        "Otros pasivos financieros, no corrientes", "PASIVO", "PNC"
    )
    assert result is not None
    assert result["standard_code"] == "PNC.05"
    assert result["method"] == "audited_statement_label"


def test_10_release_blocked_permanece_true_mientras_exista_cualquier_diferencia(tmp_path: Path):
    """10. build_gold_review_package marca release_blocked=True ante cualquier diferencia pendiente."""
    report_path = tmp_path / "report.json"
    review_output = tmp_path / "pkg"

    report_data = [{
        "file": "DOC-TEST.pdf",
        "selected_pages": [1],
        "expectations_passed": False,
        "expectation_checks": [{
            "name": "gold_mismatched_rows",
            "passed": False,
            "actual": [{
                "key": ["", "caja", 1],
                "differences": {"confidence": [0.70, 0.75]},
                "actual_line": 1,
                "expected_line": 1,
            }]
        }]
    }]
    report_path.write_text(json.dumps(report_data), encoding="utf-8")

    pkg = build_gold_review_package(report_path, review_output)
    assert pkg["summary"]["release_blocked"] is True
    assert pkg["summary"]["differences"] == 1


@pytest.mark.parametrize("invalid_conf,expected_min", [
    (float("nan"), 0.75),
    (float("inf"), 0.75),
    (float("-inf"), 0.75),
    (-0.1, 0.75),
    (1.1, 0.75),
    ("no-numerico", 0.75),
    (0.80, -0.1),
    (0.80, 1.2),
    (0.80, "invalido"),
])
def test_11_valores_no_finitos_y_fuera_de_rango_bloquean(invalid_conf, expected_min):
    """11. Confianza observada o piso no finito o fuera de [0.0, 1.0] bloquea explícitamente."""
    actual = [_base_row(confidence=invalid_conf)]
    expected = [_base_row(confidence=0.75, confidence_min=expected_min)]
    results = evaluate_gold_rows(actual, expected)
    mismatch = next(c for c in results if c["name"] == "gold_mismatched_rows")
    assert mismatch["passed"] is False
    assert len(mismatch["actual"]) == 1
    assert "confidence" in mismatch["actual"][0]["differences"]


def test_12_marcador_contrato_resumen_y_columna_cuentas_reglas(tmp_path: Path):
    """12. Validaciones fail-closed del marcador de contrato y columnas."""
    # Caso A: Columna Confianza_minima presente pero falta marcador en Resumen -> Bloquea
    wb_a = tmp_path / "caso_a.xlsx"
    with pd.ExcelWriter(wb_a, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Cuenta_original": "Caja", "Monto": 10,
            "Estado_revision": "APROBADO", "Confianza_extraccion": 0.75,
            "Confianza_minima": 0.75,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    with pytest.raises(ValueError, match="falta el marcador Gold_confidence_contract_version"):
        load_gold_rows(wb_a)

    # Caso B: Marcador presente pero falta columna Confianza_minima -> Bloquea
    wb_b = tmp_path / "caso_b.xlsx"
    with pd.ExcelWriter(wb_b, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Cuenta_original": "Caja", "Monto": 10,
            "Estado_revision": "APROBADO", "Confianza_extraccion": 0.75,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    with pytest.raises(ValueError, match="falta la columna Confianza_minima en Cuentas"):
        load_gold_rows(wb_b)

    # Caso C: Versión desconocida de contrato -> Bloquea
    wb_c = tmp_path / "caso_c.xlsx"
    with pd.ExcelWriter(wb_c, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 99,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Cuenta_original": "Caja", "Monto": 10,
            "Estado_revision": "APROBADO", "Confianza_extraccion": 0.75,
            "Confianza_minima": 0.75,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    with pytest.raises(ValueError, match="Versión de contrato de confianza no soportada: 99"):
        load_gold_rows(wb_c)


def test_13_ausencia_hojas_obligatorias_bloquea(tmp_path: Path):
    """13. La ausencia de hojas obligatorias (Resumen, Cuentas) bloquea con ValueError descriptivo."""
    # Archivo sin hoja Resumen
    no_resumen = tmp_path / "no_resumen.xlsx"
    with pd.ExcelWriter(no_resumen, engine="openpyxl") as writer:
        pd.DataFrame([{"Fila": 1}]).to_excel(writer, sheet_name="Cuentas", index=False)

    # Archivo sin hoja Cuentas
    no_cuentas = tmp_path / "no_cuentas.xlsx"
    with pd.ExcelWriter(no_cuentas, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2}]).to_excel(writer, sheet_name="Resumen", index=False)

    valid_cand = tmp_path / "valid.xlsx"
    with pd.ExcelWriter(valid_cand, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{"Fila": 1}]).to_excel(writer, sheet_name="Cuentas", index=False)

    with pytest.raises(ValueError, match="debe contener las hojas obligatorias 'Resumen' y 'Cuentas'"):
        load_gold_rows(no_resumen)

    with pytest.raises(ValueError, match="debe contener las hojas obligatorias 'Resumen' y 'Cuentas'"):
        load_gold_rows(no_cuentas)

    with pytest.raises(ValueError, match="debe contener las hojas obligatorias 'Resumen' y 'Cuentas'"):
        migrate_gold_workbook(no_resumen, valid_cand, tmp_path / "out1.xlsx")

    with pytest.raises(ValueError, match="debe contener las hojas obligatorias 'Resumen' y 'Cuentas'"):
        migrate_gold_workbook(valid_cand, no_cuentas, tmp_path / "out2.xlsx")


def test_14_gold_json_normalizacion_y_contrato(tmp_path: Path):
    """14. Gold JSON soporta normalización estricta y contrato v1 fail-closed."""
    valid_hash = "a" * 64
    # JSON con contrato v1 y piso 0.75 sin hash -> Bloquea fail-closed
    json_no_hash = tmp_path / "valid_no_hash.json"
    json_no_hash.write_text(json.dumps({
        "Gold_confidence_contract_version": 1,
        "accounts": [{
            "line": 1,
            "name": "Caja",
            "amount": 500.0,
            "Estado_revision": "APROBADO",
            "confidence": 0.75,
            "confidence_min": 0.75,
        }]
    }), encoding="utf-8")
    with pytest.raises(ValueError, match="no declara un SHA256"):
        load_gold_rows(json_no_hash)

    # JSON con contrato v1, piso 0.75 y triple hash coincidente -> Aprueba
    json_valid = tmp_path / "valid.json"
    json_valid.write_text(json.dumps({
        "Gold_confidence_contract_version": 1,
        "sha256": valid_hash,
        "accounts": [{
            "line": 1,
            "name": "Caja",
            "amount": 500.0,
            "Estado_revision": "APROBADO",
            "confidence": 0.75,
            "confidence_min": 0.75,
        }]
    }), encoding="utf-8")

    rows = load_gold_rows(
        json_valid,
        document_sha256=valid_hash,
        manifest_sha256=valid_hash,
    )
    assert len(rows) == 1
    assert rows[0]["name"] == "Caja"
    assert rows[0]["confidence_min"] == 0.75
    assert rows[0]["_authorized_sha256"] == valid_hash

    # JSON con contrato v1 pero fila sin aprobar -> Bloquea
    json_pending = tmp_path / "pending.json"
    json_pending.write_text(json.dumps({
        "Gold_confidence_contract_version": 1,
        "accounts": [{
            "line": 1,
            "name": "Caja",
            "amount": 500.0,
            "Estado_revision": "CORREGIR",
            "confidence": 0.75,
            "confidence_min": 0.75,
        }]
    }), encoding="utf-8")

    with pytest.raises(ValueError, match="Estado_revision=APROBADO o EXCLUIR"):
        load_gold_rows(json_pending)

    # JSON con columna piso pero sin marcador -> Bloquea
    json_no_marker = tmp_path / "no_marker.json"
    json_no_marker.write_text(json.dumps({
        "accounts": [{
            "line": 1,
            "name": "Caja",
            "confidence": 0.75,
            "confidence_min": 0.75,
        }]
    }), encoding="utf-8")

    with pytest.raises(ValueError, match="falta el marcador Gold_confidence_contract_version"):
        load_gold_rows(json_no_marker)


def test_15_pisos_por_fila_no_se_degradan_automaticamente(tmp_path: Path):
    """15. Demuestra pisos por fila: fila 1.0 conserva 1.0, fila 0.75 aprobada conserva 0.75, fila 1.0 no se rebaja."""
    legacy_path = tmp_path / "legacy_pisos.xlsx"
    cand_path = tmp_path / "cand_pisos.xlsx"
    out_path = tmp_path / "out_pisos.xlsx"

    # Fila 1: observada en 1.0 en legado (piso 1.0)
    # Fila 2: observada en 0.75 en legado y aprobada humanamente con piso 0.75
    legacy_rows = [
        {
            "Fila": 1,
            "Codigo_original": "10",
            "Cuenta_original": "Banco Estado",
            "Monto": 1000.0,
            "Confianza_extraccion": 1.0,
            "Confianza_minima": 1.0,
            "Estado_revision": "APROBADO",
        },
        {
            "Fila": 2,
            "Codigo_original": "20",
            "Cuenta_original": "Clientes Nacionales",
            "Monto": 2000.0,
            "Confianza_extraccion": 0.75,
            "Confianza_minima": 0.75,
            "Estado_revision": "APROBADO",
            "Observacion_analista": "Aprobado 0.75 por reconstrucción nativa",
        },
    ]

    # Candidato emite 0.75 para Fila 1 (regresión no aprobada) y 0.75 para Fila 2
    cand_rows = [
        {
            "Fila": 1,
            "Codigo_original": "10",
            "Cuenta_original": "Banco Estado",
            "Monto": 1000.0,
            "Confianza_extraccion": 0.75,  # Regresión respecto al 1.0 original
            "Estado_revision": "PENDIENTE",
        },
        {
            "Fila": 2,
            "Codigo_original": "20",
            "Cuenta_original": "Clientes Nacionales",
            "Monto": 2000.0,
            "Confianza_extraccion": 0.75,
            "Estado_revision": "PENDIENTE",
        },
    ]

    valid_hash = "b" * 64
    with pd.ExcelWriter(legacy_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
            "SHA256": valid_hash,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame(legacy_rows).to_excel(writer, sheet_name="Cuentas", index=False)

    with pd.ExcelWriter(cand_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
            "SHA256": valid_hash,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame(cand_rows).to_excel(writer, sheet_name="Cuentas", index=False)

    migrate_gold_workbook(legacy_path, cand_path, out_path)

    df_res_out = pd.read_excel(out_path, sheet_name="Resumen")
    assert df_res_out.iloc[0]["SHA256"] == valid_hash

    df_out = pd.read_excel(out_path, sheet_name="Cuentas")

    # Fila 1 (Banco Estado):
    # - Su piso NO fue rebajado a 0.75: sigue siendo 1.0
    # - Al tener confianza 0.75 bajo piso 1.0, NO se aprueba automáticamente: queda en CORREGIR
    row1 = df_out[df_out["Cuenta_original"] == "Banco Estado"].iloc[0]
    assert float(row1["Confianza_minima"]) == 1.0
    assert row1["Estado_revision"] == "CORREGIR"
    assert "bajo piso previo 1.0" in str(row1["Observacion_analista"])

    # Fila 2 (Clientes Nacionales):
    # - Su piso de 0.75 fue aprobado y preservado
    # - Al coincidir confianza 0.75 >= piso 0.75, conserva su APROBADO
    row2 = df_out[df_out["Cuenta_original"] == "Clientes Nacionales"].iloc[0]
    assert float(row2["Confianza_minima"]) == 0.75
    assert row2["Estado_revision"] == "APROBADO"


def test_16_modulo_contrato_aislado():
    """16. Validación unitaria del módulo scripts/gold_confidence_contract.py sin dependencias externas."""
    from scripts.gold_confidence_contract import (
        GOLD_CONFIDENCE_CONTRACT_VERSION,
        SUPPORTED_CONFIDENCE_CONTRACT_VERSIONS,
        is_valid_confidence,
        parse_contract_version,
        validate_confidence_floor,
        validate_candidate_confidence,
    )

    assert GOLD_CONFIDENCE_CONTRACT_VERSION == 1
    assert 1 in SUPPORTED_CONFIDENCE_CONTRACT_VERSIONS

    # is_valid_confidence
    assert is_valid_confidence(1.0) is True
    assert is_valid_confidence(0.75) is True
    assert is_valid_confidence(0.0) is True
    assert is_valid_confidence("0.75") is True
    assert is_valid_confidence(None) is False
    assert is_valid_confidence(True) is False
    assert is_valid_confidence(False) is False
    assert is_valid_confidence(-0.01) is False
    assert is_valid_confidence(1.01) is False
    assert is_valid_confidence(float("nan")) is False
    assert is_valid_confidence(float("inf")) is False
    assert is_valid_confidence("invalido") is False

    # parse_contract_version
    assert parse_contract_version(None) is None
    with pytest.raises(ValueError, match="no válida"):
        parse_contract_version("")
    with pytest.raises(ValueError, match="no válida"):
        parse_contract_version("   ")
    with pytest.raises(ValueError, match="no válida"):
        parse_contract_version(float("nan"))
    assert parse_contract_version(1) == 1
    assert parse_contract_version(1.0) == 1
    assert parse_contract_version("1") == 1

    with pytest.raises(ValueError, match="no válida"):
        parse_contract_version("invalido")
    with pytest.raises(ValueError, match="no válida"):
        parse_contract_version(1.5)
    with pytest.raises(ValueError, match="no soportada"):
        parse_contract_version(99)

    # validate_confidence_floor y validate_candidate_confidence
    assert validate_confidence_floor(0.85) == 0.85
    assert validate_candidate_confidence(0.75) == 0.75
    with pytest.raises(ValueError, match="Piso de confianza inválido"):
        validate_confidence_floor(-0.5)
    with pytest.raises(ValueError, match="Confianza de extracción candidata inválida"):
        validate_candidate_confidence("error")


def test_17_build_gold_review_package_falta_gold_bloquea(tmp_path: Path):
    """17. Hallazgo 1: Si un Gold exigido por el manifiesto no existe, build_gold_review_package lanza ValueError y nunca produce release_blocked=false."""
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(json.dumps({
        "cases": [
            {"file": "DOC-01.pdf", "gold_file": "DOC-01_gold_inexistente.xlsx"}
        ]
    }), encoding="utf-8")

    report_path = tmp_path / "report.json"
    report_path.write_text(json.dumps([
        {
            "file": "DOC-01.pdf",
            "gold_candidate": str(tmp_path / "DOC-01_cand.xlsx"),
            "expectations_passed": True,
            "expectation_checks": [],
        }
    ]), encoding="utf-8")

    cand_path = tmp_path / "DOC-01_cand.xlsx"
    with pd.ExcelWriter(cand_path, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{"Fila": 1, "Cuenta_original": "Caja", "Monto": 100}]).to_excel(writer, sheet_name="Cuentas", index=False)

    legacy_dir = tmp_path / "legacy_empty"
    legacy_dir.mkdir()

    output_dir = tmp_path / "review_package"

    with pytest.raises(ValueError, match="no existe en"):
        build_gold_review_package(
            report_path,
            output_dir,
            manifest=manifest_path,
            legacy_gold_root=legacy_dir,
        )

    # Verificar que el paquete de salida NO se haya generado o no tenga release_blocked=false
    summary_file = output_dir / "summary.json"
    if summary_file.exists():
        summary = json.loads(summary_file.read_text(encoding="utf-8"))
        assert summary.get("release_blocked") is True


def test_18_build_gold_review_package_marcador_corrupto_bloquea(tmp_path: Path):
    """18. Hallazgo 2: Marcador no numérico o corrupto en el libro legado lanza ValueError y no se ignora silenciosamente."""
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(json.dumps({
        "cases": [
            {"file": "DOC-01.pdf", "gold_file": "DOC-01_legacy.xlsx"}
        ]
    }), encoding="utf-8")

    legacy_dir = tmp_path / "legacy_dir"
    legacy_dir.mkdir()
    legacy_file = legacy_dir / "DOC-01_legacy.xlsx"

    # Libro legado con marcador de contrato corrupto (no numérico)
    with pd.ExcelWriter(legacy_file, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": "version_corrupta",
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Cuenta_original": "Caja", "Monto": 100,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    cand_path = tmp_path / "DOC-01_cand.xlsx"
    with pd.ExcelWriter(cand_path, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{"Fila": 1, "Cuenta_original": "Caja", "Monto": 100}]).to_excel(writer, sheet_name="Cuentas", index=False)

    report_path = tmp_path / "report.json"
    report_path.write_text(json.dumps([
        {
            "file": "DOC-01.pdf",
            "gold_candidate": str(cand_path),
            "expectations_passed": True,
            "expectation_checks": [],
        }
    ]), encoding="utf-8")

    output_dir = tmp_path / "review_package"

    with pytest.raises(ValueError, match="no válida"):
        build_gold_review_package(
            report_path,
            output_dir,
            manifest=manifest_path,
            legacy_gold_root=legacy_dir,
        )


def test_19_preservacion_hojas_auxiliares_migracion(tmp_path: Path):
    """19. Hallazgo 3: La migración preserva hojas auxiliares del candidato (e.g. Controles, Notas) sin descartarlas."""
    import openpyxl

    legacy_path = tmp_path / "legacy.xlsx"
    with pd.ExcelWriter(legacy_path, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Estado_revision": "APROBADO",
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    cand_path = tmp_path / "cand_con_hojas.xlsx"
    with pd.ExcelWriter(cand_path, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Confianza_extraccion": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)
        pd.DataFrame([{"Control": "Total Activos", "Valor": 100.0}]).to_excel(writer, sheet_name="Controles", index=False)
        pd.DataFrame([{"Nota": "Auditoría limpia"}]).to_excel(writer, sheet_name="Notas", index=False)

    out_path = tmp_path / "migrated_preserved.xlsx"
    migrate_gold_workbook(legacy_path, cand_path, out_path)

    wb = openpyxl.load_workbook(out_path)
    assert "Resumen" in wb.sheetnames
    assert "Cuentas" in wb.sheetnames
    assert "Controles" in wb.sheetnames
    assert "Notas" in wb.sheetnames

    # Verificar que los datos de las hojas auxiliares estén intactos
    df_controles = pd.read_excel(out_path, sheet_name="Controles")
    assert df_controles.iloc[0]["Control"] == "Total Activos"
    assert df_controles.iloc[0]["Valor"] == 100.0


def test_20_validacion_estricta_piso_y_candidato_en_migracion(tmp_path: Path):
    """20. Hallazgo 4: Validación fail-closed de candidate_conf y legacy_min en migrate_gold_workbook."""
    # Caso A: Confianza candidata no válida (negativa o string) -> ValueError
    legacy_path = tmp_path / "leg_ok.xlsx"
    with pd.ExcelWriter(legacy_path, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Estado_revision": "APROBADO", "Confianza_minima": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    cand_bad = tmp_path / "cand_bad.xlsx"
    with pd.ExcelWriter(cand_bad, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Confianza_extraccion": -0.5,  # Inválido
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    with pytest.raises(ValueError, match="candidata tiene Confianza_extraccion inválida"):
        migrate_gold_workbook(legacy_path, cand_bad, tmp_path / "out_bad.xlsx")

    # Caso B: Piso de confianza en libro legado no válido -> ValueError
    legacy_bad = tmp_path / "leg_bad.xlsx"
    with pd.ExcelWriter(legacy_bad, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Estado_revision": "APROBADO", "Confianza_minima": "no_numerica",  # Inválido
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    cand_ok = tmp_path / "cand_ok.xlsx"
    with pd.ExcelWriter(cand_ok, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Confianza_extraccion": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    with pytest.raises(ValueError, match="libro legado tiene piso de confianza inválido"):
        migrate_gold_workbook(legacy_bad, cand_ok, tmp_path / "out_bad2.xlsx")

    # Caso C: Libro legado carece por completo de columnas de piso (contrato legado pre-schema 2) -> adopta 1.0
    legacy_pre_s2 = tmp_path / "leg_pre_s2.xlsx"
    with pd.ExcelWriter(legacy_pre_s2, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 1}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Estado_revision": "APROBADO",
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    out_pre_s2 = tmp_path / "out_pre_s2.xlsx"
    migrate_gold_workbook(legacy_pre_s2, cand_ok, out_pre_s2)
    df_pre = pd.read_excel(out_pre_s2, sheet_name="Cuentas")
    assert float(df_pre.iloc[0]["Confianza_minima"]) == 1.0
    assert df_pre.iloc[0]["Estado_revision"] == "APROBADO"


def test_21_marcador_corrupto_en_legacy_bloquea_migracion_sin_archivo_salida(tmp_path: Path):
    """21. Marcador corrupto en libro legado bloquea migrate_gold_workbook() con ValueError y no deja archivo de salida."""
    legacy_path = tmp_path / "leg_corrupt.xlsx"
    with pd.ExcelWriter(legacy_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": "corrupto",
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Estado_revision": "APROBADO", "Confianza_minima": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    cand_path = tmp_path / "cand_ok.xlsx"
    with pd.ExcelWriter(cand_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Confianza_extraccion": 1.0, "Confianza_minima": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    out_path = tmp_path / "salida_no_debe_existir.xlsx"
    with pytest.raises(ValueError, match="no válida"):
        migrate_gold_workbook(legacy_path, cand_path, out_path)

    assert not out_path.exists(), "Tras el bloqueo no debe quedar un archivo de salida utilizable."


def test_22_marcador_corrupto_en_candidato_bloquea_migracion_sin_archivo_salida(tmp_path: Path):
    """22. Marcador corrupto en candidato bloquea migrate_gold_workbook() con ValueError y no lo normaliza silenciosamente."""
    legacy_path = tmp_path / "leg_ok.xlsx"
    with pd.ExcelWriter(legacy_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Estado_revision": "APROBADO", "Confianza_minima": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    cand_corrupt = tmp_path / "cand_corrupt.xlsx"
    with pd.ExcelWriter(cand_corrupt, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": "corrupto_no_normalizable",
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Confianza_extraccion": 1.0, "Confianza_minima": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    out_path = tmp_path / "salida_cand_corrupto.xlsx"
    with pytest.raises(ValueError, match="no válida"):
        migrate_gold_workbook(legacy_path, cand_corrupt, out_path)

    assert not out_path.exists(), "Tras el bloqueo no debe quedar un archivo de salida utilizable."


def test_23_marcador_version_desconocida_bloquea_migracion(tmp_path: Path):
    """23. Versión desconocida (e.g. 99) en libro legado o candidato bloquea con ValueError y sin archivo de salida."""
    legacy_v99 = tmp_path / "leg_v99.xlsx"
    with pd.ExcelWriter(legacy_v99, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 99,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Estado_revision": "APROBADO", "Confianza_minima": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    cand_path = tmp_path / "cand_ok.xlsx"
    with pd.ExcelWriter(cand_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Confianza_extraccion": 1.0, "Confianza_minima": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    out_path = tmp_path / "salida_v99.xlsx"
    with pytest.raises(ValueError, match="no soportada: 99"):
        migrate_gold_workbook(legacy_v99, cand_path, out_path)

    assert not out_path.exists(), "Tras el bloqueo no debe quedar archivo de salida."


def test_24_marcador_ausente_en_legado_valido_conserva_comportamiento_compatible(tmp_path: Path):
    """24. Marcador ausente en libro legado válido pre-schema 2 conserva el comportamiento compatible y asigna piso 1.0."""
    legacy_legacy = tmp_path / "leg_pre_s2_valido.xlsx"
    with pd.ExcelWriter(legacy_legacy, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 1,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Estado_revision": "APROBADO",
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    cand_path = tmp_path / "cand_para_legado.xlsx"
    with pd.ExcelWriter(cand_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Codigo_original": "1", "Cuenta_original": "Caja",
            "Monto": 100.0, "Confianza_extraccion": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    out_path = tmp_path / "salida_legado_compatible.xlsx"
    res = migrate_gold_workbook(legacy_legacy, cand_path, out_path)
    assert res == out_path
    assert out_path.exists()

    df_res = pd.read_excel(out_path, sheet_name="Resumen")
    assert int(df_res.iloc[0]["Gold_confidence_contract_version"]) == 1
    assert int(df_res.iloc[0]["Gold_schema_version"]) == 2

    df_cue = pd.read_excel(out_path, sheet_name="Cuentas")
    assert float(df_cue.iloc[0]["Confianza_minima"]) == 1.0
    assert df_cue.iloc[0]["Estado_revision"] == "APROBADO"


def test_25_vinculacion_piso_al_hash_documental_y_bloqueo_por_alteracion(tmp_path: Path):
    """25. Verifica que el piso 0.75 esté estrictamente vinculado al hash de DOC-01 y cualquier alteración bloquee."""
    from scripts.certify_local_corpus import evaluate_expectations

    hash_autorizado_doc01 = "2c51270ea91ba1e805f99a745c217356312bf8b17d2582f40df5d6e337f30711"

    # Caso A: DOC-01 con hash autorizado pasa la verificación de hash
    result_ok = {
        "file": "DOC-01.pdf",
        "sha256": hash_autorizado_doc01,
        "raw_accounts": 37,
        "qualified_accounts": 37,
        "detected_periods": ["2021", "2022"],
        "detected_currencies": ["CLP"],
    }
    expectations = {
        "sha256": hash_autorizado_doc01,
        "exact_raw_accounts": 37,
        "exact_qualified_accounts": 37,
        "periods": ["2021", "2022"],
        "currencies": ["CLP"],
    }
    checks_ok, passed_ok = evaluate_expectations(result_ok, expectations)
    assert passed_ok is True
    sha_check = next(c for c in checks_ok if c["name"] == "sha256")
    assert sha_check["passed"] is True
    assert sha_check["actual"] == hash_autorizado_doc01

    # Caso B: Alteración del hash en documento o expectativa bloquea inmediatamente
    hash_alterado = "ffffffffffffffffffffffffffffffffffffffffffffffffffffffffffffffff"
    result_tampered = dict(result_ok, sha256=hash_alterado)
    checks_tampered, passed_tampered = evaluate_expectations(result_tampered, expectations)
    assert passed_tampered is False
    sha_check_tampered = next(c for c in checks_tampered if c["name"] == "sha256")
    assert sha_check_tampered["passed"] is False
    assert sha_check_tampered["actual"] != sha_check_tampered["expected"]

    # Caso C: Verificar que en el paquete de revisión una expectativa fallida produce release_blocked=True
    report_file = tmp_path / "rep_tampered.json"
    report_file.write_text(json.dumps([{
        "file": "DOC-01.pdf",
        "gold_candidate": str(tmp_path / "cand.xlsx"),
        "expectations_passed": False,  # Falló por hash
        "expectation_checks": checks_tampered,
    }]), encoding="utf-8")

    cand_file = tmp_path / "cand.xlsx"
    with pd.ExcelWriter(cand_file, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2, "Gold_confidence_contract_version": 1}]).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        pd.DataFrame([{"Fila": 1, "Cuenta_original": "Caja", "Monto": 100.0}]).to_excel(
            writer, sheet_name="Cuentas", index=False
        )

    out_review = tmp_path / "review_blocked_by_hash"
    build_gold_review_package(report_file, out_review)
    summary = json.loads((out_review / "summary.json").read_text(encoding="utf-8"))
    assert summary["release_blocked"] is True, "Cualquier alteración de hash debe mantener release_blocked=True."


def test_26_preservacion_salida_preexistente_ante_fallo_marcador_corrupto(tmp_path: Path):
    """26. Demuestra que migrate_gold_workbook() preserva la salida preexistente ante fallo por marcador corrupto."""
    out_path = tmp_path / "target_output.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        pd.DataFrame([{"info": "preexistente_valido", "val": 12345}]).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        pd.DataFrame([{"Fila": 1, "Cuenta": "Caja"}]).to_excel(
            writer, sheet_name="Cuentas", index=False
        )
    sha_pre = hashlib.sha256(out_path.read_bytes()).hexdigest()

    legacy_valid = tmp_path / "legacy_ok.xlsx"
    with pd.ExcelWriter(legacy_valid, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2, "Gold_confidence_contract_version": 1}]).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        pd.DataFrame([{"Fila": 1, "Cuenta_original": "Caja", "Monto": 100.0, "Confianza_extraccion": 1.0, "Confianza_minima": 1.0, "Estado_revision": "APROBADO"}]).to_excel(
            writer, sheet_name="Cuentas", index=False
        )

    cand_corrupt = tmp_path / "cand_corrupt.xlsx"
    with pd.ExcelWriter(cand_corrupt, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2, "Gold_confidence_contract_version": "corrupto"}]).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        pd.DataFrame([{"Fila": 1, "Cuenta_original": "Caja", "Monto": 100.0, "Confianza_extraccion": 1.0}]).to_excel(
            writer, sheet_name="Cuentas", index=False
        )

    with pytest.raises(ValueError, match="contrato de confianza no válida"):
        migrate_gold_workbook(legacy_valid, cand_corrupt, out_path)

    assert out_path.exists(), "La salida preexistente debe seguir existiendo."
    sha_post = hashlib.sha256(out_path.read_bytes()).hexdigest()
    assert sha_post == sha_pre, "El SHA-256 de la salida preexistente debe permanecer idéntico tras el fallo."

    # Verificar que no queden temporales huérfanos
    tmps = list(tmp_path.glob("*.tmp_*")) + list(tmp_path.glob(".*.tmp_*"))
    assert len(tmps) == 0, f"No deben quedar archivos temporales huérfanos: {tmps}"


def test_27_salida_inexistente_mas_marcador_corrupto_no_crea_salida(tmp_path: Path):
    """27. Si la salida no preexistía y ocurre fallo por marcador corrupto, no se crea output."""
    out_path = tmp_path / "nonexistent_output.xlsx"
    assert not out_path.exists()

    legacy_valid = tmp_path / "legacy_27.xlsx"
    with pd.ExcelWriter(legacy_valid, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2, "Gold_confidence_contract_version": 1}]).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        pd.DataFrame([{"Fila": 1, "Cuenta_original": "Caja", "Monto": 100.0, "Confianza_extraccion": 1.0, "Confianza_minima": 1.0, "Estado_revision": "APROBADO"}]).to_excel(
            writer, sheet_name="Cuentas", index=False
        )

    cand_corrupt = tmp_path / "cand_27.xlsx"
    with pd.ExcelWriter(cand_corrupt, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2, "Gold_confidence_contract_version": None}]).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        pd.DataFrame([{"Fila": 1, "Cuenta_original": "Caja", "Monto": 100.0, "Confianza_extraccion": 1.0}]).to_excel(
            writer, sheet_name="Cuentas", index=False
        )

    with pytest.raises(ValueError):
        migrate_gold_workbook(legacy_valid, cand_corrupt, out_path)

    assert not out_path.exists(), "No debe haberse creado el archivo de salida."
    tmps = list(tmp_path.glob("*.tmp_*")) + list(tmp_path.glob(".*.tmp_*"))
    assert len(tmps) == 0


def test_28_fallo_durante_escritura_o_validacion_conserva_salida_anterior(tmp_path: Path, monkeypatch):
    """28. Fallo en validación o escritura del temporal conserva intacta la salida anterior."""
    out_path = tmp_path / "target_out28.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        pd.DataFrame([{"preexistente": 999}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{"Fila": 1}]).to_excel(writer, sheet_name="Cuentas", index=False)
    sha_pre = hashlib.sha256(out_path.read_bytes()).hexdigest()

    legacy_valid = tmp_path / "legacy_28.xlsx"
    with pd.ExcelWriter(legacy_valid, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2, "Gold_confidence_contract_version": 1}]).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        pd.DataFrame([{"Fila": 1, "Cuenta_original": "Caja", "Monto": 100.0, "Confianza_extraccion": 1.0, "Confianza_minima": 1.0, "Estado_revision": "APROBADO"}]).to_excel(
            writer, sheet_name="Cuentas", index=False
        )

    cand_valid = tmp_path / "cand_28.xlsx"
    with pd.ExcelWriter(cand_valid, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2, "Gold_confidence_contract_version": 1}]).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        pd.DataFrame([{"Fila": 1, "Cuenta_original": "Caja", "Monto": 100.0, "Confianza_extraccion": 1.0}]).to_excel(
            writer, sheet_name="Cuentas", index=False
        )

    # Simular fallo durante validate_workbook_resumen_and_cuentas
    import scripts.migrate_gold_workbook as mgw
    orig_validate = mgw.validate_workbook_resumen_and_cuentas
    def mock_validate(path, label="libro", enforce_confidence_coherence=False):
        if label == "temporal":
            raise RuntimeError("Fallo simulado en validación del archivo temporal")
        return orig_validate(path, label=label, enforce_confidence_coherence=enforce_confidence_coherence)
    monkeypatch.setattr(mgw, "validate_workbook_resumen_and_cuentas", mock_validate)

    with pytest.raises(RuntimeError, match="Fallo simulado"):
        migrate_gold_workbook(legacy_valid, cand_valid, out_path)

    assert out_path.exists()
    assert hashlib.sha256(out_path.read_bytes()).hexdigest() == sha_pre
    tmps = list(tmp_path.glob("*.tmp_*")) + list(tmp_path.glob(".*.tmp_*"))
    assert len(tmps) == 0


def test_29_migracion_exitosa_reemplaza_atomicamente_salida_anterior_y_preserva_hojas(tmp_path: Path):
    """29. Migración exitosa reemplaza atómicamente la salida previa y preserva hojas auxiliares."""
    out_path = tmp_path / "target_out29.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        pd.DataFrame([{"version_vieja": 1}]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{"Fila": 1}]).to_excel(writer, sheet_name="Cuentas", index=False)
    sha_pre = hashlib.sha256(out_path.read_bytes()).hexdigest()

    legacy_valid = tmp_path / "legacy_29.xlsx"
    with pd.ExcelWriter(legacy_valid, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2, "Gold_confidence_contract_version": 1}]).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        pd.DataFrame([{"Fila": 1, "Cuenta_original": "Caja", "Monto": 100.0, "Confianza_extraccion": 1.0, "Confianza_minima": 1.0, "Estado_revision": "APROBADO"}]).to_excel(
            writer, sheet_name="Cuentas", index=False
        )

    cand_valid = tmp_path / "cand_29.xlsx"
    with pd.ExcelWriter(cand_valid, engine="openpyxl") as writer:
        pd.DataFrame([{"Gold_schema_version": 2, "Gold_confidence_contract_version": 1}]).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        pd.DataFrame([{"Fila": 1, "Cuenta_original": "Caja", "Monto": 100.0, "Confianza_extraccion": 1.0}]).to_excel(
            writer, sheet_name="Cuentas", index=False
        )
        pd.DataFrame([{"Control": "Total_Activo", "Cumple": True}]).to_excel(
            writer, sheet_name="Controles", index=False
        )
        pd.DataFrame([{"Advertencia": "Ninguna"}]).to_excel(
            writer, sheet_name="Advertencias", index=False
        )

    res = migrate_gold_workbook(legacy_valid, cand_valid, out_path)
    assert res == out_path
    assert out_path.exists()
    sha_post = hashlib.sha256(out_path.read_bytes()).hexdigest()
    assert sha_post != sha_pre, "El contenido debió ser actualizado atómicamente."

    excel = pd.ExcelFile(out_path)
    assert set(excel.sheet_names) == {"Resumen", "Cuentas", "Controles", "Advertencias"}
    df_res = pd.read_excel(excel, sheet_name="Resumen")
    assert int(df_res.iloc[0]["Gold_schema_version"]) == 2
    assert int(df_res.iloc[0]["Gold_confidence_contract_version"]) == 1
    df_cue = pd.read_excel(excel, sheet_name="Cuentas")
    assert float(df_cue.iloc[0]["Confianza_minima"]) == 1.0


def test_30_piso_sub_uno_triple_hash_binding_aprueba(tmp_path: Path):
    """30. Piso 0.75 con hash autorizado, hash real y manifiesto coincidentes aprueba."""
    auth_hash = "2c51270ea91ba1e805f99a745c217356312bf8b17d2582f40df5d6e337f30711"
    gold_file = tmp_path / "gold_doc01.xlsx"
    base_acc = _base_row(account_code="1", name="Caja", amount=100.0, confidence=0.75)
    with pd.ExcelWriter(gold_file, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
            "SHA256": auth_hash,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1,
            "Codigo_original": "1",
            "Cuenta_original": "Caja",
            "Jerarquia_contable": base_acc["accounting_hierarchy"],
            "Origen": base_acc["origin"],
            "Monto": 100.0,
            "Montos_periodos_JSON": json.dumps(base_acc["period_amounts"]),
            "Montos_columnas_JSON": json.dumps(base_acc["column_amounts"]),
            "Codigo_homologado": base_acc["standard_code"],
            "Metodo_clasificacion": base_acc["classification_method"],
            "Requiere_revision": base_acc["requires_review"],
            "Es_control_total": base_acc["is_total"],
            "Confianza_extraccion": 0.75,
            "Confianza_minima": 0.75,
            "Estado_revision": "APROBADO",
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    rows = load_gold_rows(
        gold_file,
        document_sha256=auth_hash,
        manifest_sha256=auth_hash,
    )
    assert len(rows) == 1
    assert rows[0]["confidence_min"] == 0.75
    assert rows[0]["_authorized_sha256"] == auth_hash

    # Evaluar expectativas con documento real coincidente
    result = {
        "file": "DOC-01.pdf",
        "sha256": auth_hash,
        "accounts": [base_acc],
    }
    expectations = {
        "sha256": auth_hash,
        "_gold_rows": rows,
    }
    checks, passed = evaluate_expectations(result, expectations)
    assert passed is True
    assert all(c["passed"] for c in checks)


def test_31_piso_sub_uno_sin_hash_en_gold_bloquea(tmp_path: Path):
    """31. Piso 0.75 sin hash en Resumen del libro Gold bloquea fail-closed con ValueError."""
    valid_hash = "2c51270ea91ba1e805f99a745c217356312bf8b17d2582f40df5d6e337f30711"
    gold_no_hash = tmp_path / "gold_nohash.xlsx"
    with pd.ExcelWriter(gold_no_hash, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
            # Falta SHA256
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1,
            "Codigo_original": "1",
            "Cuenta_original": "Caja",
            "Monto": 100.0,
            "Confianza_extraccion": 0.75,
            "Confianza_minima": 0.75,
            "Estado_revision": "APROBADO",
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    with pytest.raises(ValueError, match="no declara un SHA256 documental autorizado"):
        load_gold_rows(
            gold_no_hash,
            document_sha256=valid_hash,
            manifest_sha256=valid_hash,
        )


def test_32_piso_sub_uno_sin_expectativa_sha256_en_manifiesto_bloquea(tmp_path: Path):
    """32. Piso 0.75 sin expectativa sha256 en manifiesto bloquea fail-closed (FLOOR_BLOCKED_WITHOUT_HASH_EXPECTATION)."""
    valid_hash = "2c51270ea91ba1e805f99a745c217356312bf8b17d2582f40df5d6e337f30711"
    gold_file = tmp_path / "gold_32.xlsx"
    with pd.ExcelWriter(gold_file, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
            "SHA256": valid_hash,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1,
            "Codigo_original": "1",
            "Cuenta_original": "Caja",
            "Monto": 100.0,
            "Confianza_extraccion": 0.75,
            "Confianza_minima": 0.75,
            "Estado_revision": "APROBADO",
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    # 1. En load_gold_rows, si manifest_sha256 falta -> Bloquea
    with pytest.raises(ValueError, match="manifiesto no declara la expectativa sha256"):
        load_gold_rows(
            gold_file,
            document_sha256=valid_hash,
            manifest_sha256=None,
        )

    # 2. En evaluate_expectations, si expectations omite sha256 -> Bloquea
    rows = [{
        "line": 1,
        "account_code": "1",
        "name": "Caja",
        "amount": 100.0,
        "confidence": 0.75,
        "confidence_min": 0.75,
        "_authorized_sha256": valid_hash,
    }]
    result = {
        "file": "DOC-01.pdf",
        "sha256": valid_hash,
        "accounts": rows,
    }
    expectations_no_sha = {
        "_gold_rows": rows,  # Sin 'sha256'
    }
    with pytest.raises(ValueError, match="manifiesto no declara la expectativa sha256"):
        evaluate_expectations(result, expectations_no_sha)


def test_33_piso_sub_uno_hash_malformado_o_discrepante_bloquea(tmp_path: Path):
    """33. Hash malformado, discrepante entre Gold y manifiesto, o documento distinto bloquea."""
    auth_hash = "2c51270ea91ba1e805f99a745c217356312bf8b17d2582f40df5d6e337f30711"
    wrong_hash = "ffffffffffffffffffffffffffffffffffffffffffffffffffffffffffffffff"
    short_hash = "2c51270ea91ba1e8"

    # A. Hash malformado en Gold
    gold_malformed = tmp_path / "gold_malformed.xlsx"
    with pd.ExcelWriter(gold_malformed, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
            "SHA256": short_hash,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1,
            "Cuenta_original": "Caja",
            "Monto": 100.0,
            "Confianza_extraccion": 0.75,
            "Confianza_minima": 0.75,
            "Estado_revision": "APROBADO",
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    with pytest.raises(ValueError, match="malformado o inválido"):
        load_gold_rows(gold_malformed, document_sha256=auth_hash, manifest_sha256=auth_hash)

    # B. Gold correcto pero discrepancia con manifiesto
    gold_ok = tmp_path / "gold_ok.xlsx"
    with pd.ExcelWriter(gold_ok, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
            "SHA256": auth_hash,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1,
            "Cuenta_original": "Caja",
            "Monto": 100.0,
            "Confianza_extraccion": 0.75,
            "Confianza_minima": 0.75,
            "Estado_revision": "APROBADO",
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    with pytest.raises(ValueError, match="Discrepancia de hash"):
        load_gold_rows(gold_ok, document_sha256=auth_hash, manifest_sha256=wrong_hash)

    # C. Mismo nombre de archivo pero documento real con contenido / hash distinto
    with pytest.raises(ValueError, match="Discrepancia de hash"):
        load_gold_rows(gold_ok, document_sha256=wrong_hash, manifest_sha256=auth_hash)


def test_34_piso_uno_mantiene_igualdad_estricta_sin_requerir_hash_binding(tmp_path: Path):
    """34. Fila con piso 1.0 mantiene igualdad estricta sin exigir vinculación a hash."""
    gold_piso_uno = tmp_path / "gold_piso1.xlsx"
    acc_base = _base_row(account_code="1", name="Caja", amount=100.0, confidence=1.0)
    with pd.ExcelWriter(gold_piso_uno, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2,
            "Gold_confidence_contract_version": 1,
            # No declara SHA256 (no obligatorio para piso 1.0)
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1,
            "Codigo_original": "1",
            "Cuenta_original": "Caja",
            "Jerarquia_contable": acc_base["accounting_hierarchy"],
            "Origen": acc_base["origin"],
            "Monto": 100.0,
            "Montos_periodos_JSON": json.dumps(acc_base["period_amounts"]),
            "Montos_columnas_JSON": json.dumps(acc_base["column_amounts"]),
            "Codigo_homologado": acc_base["standard_code"],
            "Metodo_clasificacion": acc_base["classification_method"],
            "Requiere_revision": acc_base["requires_review"],
            "Es_control_total": acc_base["is_total"],
            "Confianza_extraccion": 1.0,
            "Confianza_minima": 1.0,
            "Estado_revision": "APROBADO",
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    # Carga sin hash pasa porque el piso es 1.0
    rows = load_gold_rows(gold_piso_uno)
    assert len(rows) == 1
    assert rows[0]["confidence_min"] == 1.0

    # Si el candidato emite 0.75 para piso 1.0 -> Falla
    result_regresion = {
        "accounts": [_base_row(account_code="1", name="Caja", amount=100.0, confidence=0.75)]
    }
    checks_bad, passed_bad = evaluate_expectations(result_regresion, {"_gold_rows": rows})
    assert passed_bad is False

    # Si el candidato emite 1.0 para piso 1.0 -> Aprueba
    result_ok = {
        "accounts": [acc_base]
    }
    checks_ok, passed_ok = evaluate_expectations(result_ok, {"_gold_rows": rows})
    assert passed_ok is True


def test_35_doc02_y_doc03_mantienen_piso_uno_estricto_y_no_se_rebajan():
    """35. DOC-02 y DOC-03 no tienen autorización de rebaja y mantienen piso 1.0 estricto."""
    # Fila simulada para DOC-02 o DOC-03
    doc02_gold_row = {
        "line": 1,
        "account_code": "1101",
        "name": "Banco Santander",
        "amount": 5000.0,
        "confidence": 1.0,
        "confidence_min": 1.0,
    }

    # Si el motor extrae 0.75, debe fallar (no se beneficia del 0.75 de DOC-01)
    actual_doc02_regresion = {
        "line": 1,
        "account_code": "1101",
        "name": "Banco Santander",
        "amount": 5000.0,
        "confidence": 0.75,
    }
    checks = evaluate_gold_rows([actual_doc02_regresion], [doc02_gold_row], tolerance=0.01)
    assert not all(c["passed"] for c in checks)
    mismatches = next(c for c in checks if c["name"] == "gold_mismatched_rows")
    assert "confidence" in mismatches["actual"][0]["differences"]
    diff = mismatches["actual"][0]["differences"]["confidence"]
    assert diff == [0.75, 1.0]


def test_36_manual_gold_rows_subone_without_authorized_sha256_blocked():
    """36. Filas Gold construidas manualmente con piso < 1.0 sin _authorized_sha256 bloquean con ValueError."""
    auth_hash = "2c51270ea91ba1e805f99a745c217356312bf8b17d2582f40df5d6e337f30711"
    row_manual = _base_row(account_code="1", name="Caja", amount=100.0, confidence=0.75, confidence_min=0.75)
    result = {
        "file": "DOC-01.pdf",
        "sha256": auth_hash,
        "accounts": [row_manual],
    }
    expectations = {
        "sha256": auth_hash,
        "_gold_rows": [row_manual],
    }
    with pytest.raises(ValueError, match="no declara un hash documental autorizado"):
        evaluate_expectations(result, expectations)


def test_37_manifest_sha_cannot_substitute_gold_authorized_sha():
    """37. manifest_sha presente no puede sustituir la ausencia de hash autorizado en el Gold."""
    manifest_sha = "2c51270ea91ba1e805f99a745c217356312bf8b17d2582f40df5d6e337f30711"
    row_manual = _base_row(account_code="1", name="Caja", amount=100.0, confidence=0.75, confidence_min=0.75)
    row_manual["_authorized_sha256"] = None

    result = {
        "file": "DOC-01.pdf",
        "sha256": manifest_sha,
        "accounts": [row_manual],
    }
    expectations = {
        "sha256": manifest_sha,
        "_gold_rows": [row_manual],
    }
    with pytest.raises(ValueError, match="no declara un hash documental autorizado"):
        evaluate_expectations(result, expectations)


def test_38_two_gold_rows_with_different_authorized_sha_blocked():
    """38. Filas con piso < 1.0 que declaran hashes autorizados discrepantes entre sí bloquean."""
    hash1 = "2c51270ea91ba1e805f99a745c217356312bf8b17d2582f40df5d6e337f30711"
    hash2 = "1111111111111111111111111111111111111111111111111111111111111111"
    row1 = _base_row(line=1, account_code="1", name="Caja", amount=100.0, confidence=0.75, confidence_min=0.75, _authorized_sha256=hash1)
    row2 = _base_row(line=2, account_code="2", name="Banco", amount=200.0, confidence=0.75, confidence_min=0.75, _authorized_sha256=hash2)

    result = {
        "file": "DOC-01.pdf",
        "sha256": hash1,
        "accounts": [row1, row2],
    }
    expectations = {
        "sha256": hash1,
        "_gold_rows": [row1, row2],
    }
    with pytest.raises(ValueError, match="múltiples hashes autorizados distintos"):
        evaluate_expectations(result, expectations)


def test_39_malformed_authorized_sha_in_manual_rows_blocked():
    """39. Hash autorizado malformado en fila con piso < 1.0 bloquea con ValueError."""
    bad_hash = "not_a_valid_sha256_hash_12345"
    row = _base_row(line=1, account_code="1", name="Caja", amount=100.0, confidence=0.75, confidence_min=0.75, _authorized_sha256=bad_hash)

    result = {
        "file": "DOC-01.pdf",
        "sha256": bad_hash,
        "accounts": [row],
    }
    expectations = {
        "sha256": bad_hash,
        "_gold_rows": [row],
    }
    with pytest.raises(ValueError, match="malformado o inválido"):
        evaluate_expectations(result, expectations)


def test_40_migration_with_subone_and_no_hash_blocked_all_scenarios(tmp_path: Path):
    """40. Migración con piso < 1.0 bloquea en todas las combinaciones inválidas y preserva salida preexistente."""
    valid_hash = "2c51270ea91ba1e805f99a745c217356312bf8b17d2582f40df5d6e337f30711"
    diff_hash = "3333333333333333333333333333333333333333333333333333333333333333"
    bad_hash = "abcd1234"

    scenarios = [
        (None, None, "exige validación de hash pero no declara un SHA-256"),
        (valid_hash, None, "exige validación de hash pero no declara un SHA-256"),
        (None, valid_hash, "exige validación de hash pero no declara un SHA-256"),
        (valid_hash, diff_hash, "Discrepancia de SHA-256"),
        (bad_hash, bad_hash, "malformado o inválido"),
    ]

    for idx, (leg_h, cand_h, err_regex) in enumerate(scenarios):
        scen_dir = tmp_path / f"scen_{idx}"
        scen_dir.mkdir()
        legacy_path = scen_dir / "legacy.xlsx"
        cand_path = scen_dir / "cand.xlsx"
        out_path = scen_dir / "out.xlsx"

        sentinel_content = f"PREEXISTING_VALID_OUTPUT_{idx}".encode("utf-8")
        out_path.write_bytes(sentinel_content)
        orig_hash = hashlib.sha256(sentinel_content).hexdigest()

        leg_summary = {"Gold_schema_version": 2, "Gold_confidence_contract_version": 1}
        if leg_h:
            leg_summary["SHA256"] = leg_h
        with pd.ExcelWriter(legacy_path, engine="openpyxl") as writer:
            pd.DataFrame([leg_summary]).to_excel(writer, sheet_name="Resumen", index=False)
            pd.DataFrame([{
                "Fila": 1, "Cuenta_original": "Caja", "Monto": 100,
                "Estado_revision": "APROBADO", "Confianza_extraccion": 0.75, "Confianza_minima": 0.75,
            }]).to_excel(writer, sheet_name="Cuentas", index=False)

        cand_summary = {"Gold_schema_version": 2, "Gold_confidence_contract_version": 1}
        if cand_h:
            cand_summary["SHA256"] = cand_h
        with pd.ExcelWriter(cand_path, engine="openpyxl") as writer:
            pd.DataFrame([cand_summary]).to_excel(writer, sheet_name="Resumen", index=False)
            pd.DataFrame([{
                "Fila": 1, "Cuenta_original": "Caja", "Monto": 100,
                "Estado_revision": "APROBADO", "Confianza_extraccion": 0.75, "Confianza_minima": 0.75,
            }]).to_excel(writer, sheet_name="Cuentas", index=False)

        with pytest.raises(ValueError, match=err_regex):
            migrate_gold_workbook(legacy_path, cand_path, out_path)

        assert out_path.exists()
        current_hash = hashlib.sha256(out_path.read_bytes()).hexdigest()
        assert current_hash == orig_hash, f"Salida preexistente alterada en escenario {idx}"


def test_41_migration_subone_valid_sha256_preserves_sha256_in_output(tmp_path: Path):
    """41. Migración con piso < 1.0 y SHA256 válido coincidente tiene éxito y preserva SHA256 en salida."""
    valid_hash = "2c51270ea91ba1e805f99a745c217356312bf8b17d2582f40df5d6e337f30711"
    legacy_path = tmp_path / "legacy.xlsx"
    cand_path = tmp_path / "cand.xlsx"
    out_path = tmp_path / "out.xlsx"

    with pd.ExcelWriter(legacy_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2, "Gold_confidence_contract_version": 1, "SHA256": valid_hash,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Cuenta_original": "Caja", "Origen": "activo", "Monto": 100,
            "Estado_revision": "APROBADO", "Confianza_extraccion": 0.75, "Confianza_minima": 0.75,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    with pd.ExcelWriter(cand_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2, "Gold_confidence_contract_version": 1, "SHA256": valid_hash,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Cuenta_original": "Caja", "Origen": "activo", "Monto": 100,
            "Estado_revision": "PENDIENTE", "Confianza_extraccion": 0.75, "Confianza_minima": 0.75,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    migrate_gold_workbook(legacy_path, cand_path, out_path)
    assert out_path.exists()
    df_res = pd.read_excel(out_path, sheet_name="Resumen")
    assert "SHA256" in df_res.columns
    assert df_res["SHA256"].iloc[0] == valid_hash


def test_42_migration_floor_one_preserves_compatibility_without_sha(tmp_path: Path):
    """42. Migración con piso 1.0 no exige SHA256 y mantiene retrocompatibilidad."""
    legacy_path = tmp_path / "legacy_piso1.xlsx"
    cand_path = tmp_path / "cand_piso1.xlsx"
    out_path = tmp_path / "out_piso1.xlsx"

    with pd.ExcelWriter(legacy_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2, "Gold_confidence_contract_version": 1,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Cuenta_original": "Caja", "Origen": "activo", "Monto": 100,
            "Estado_revision": "APROBADO", "Confianza_extraccion": 1.0, "Confianza_minima": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    with pd.ExcelWriter(cand_path, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Gold_schema_version": 2, "Gold_confidence_contract_version": 1,
        }]).to_excel(writer, sheet_name="Resumen", index=False)
        pd.DataFrame([{
            "Fila": 1, "Cuenta_original": "Caja", "Origen": "activo", "Monto": 100,
            "Estado_revision": "PENDIENTE", "Confianza_extraccion": 1.0, "Confianza_minima": 1.0,
        }]).to_excel(writer, sheet_name="Cuentas", index=False)

    migrate_gold_workbook(legacy_path, cand_path, out_path)
    assert out_path.exists()
    df_cuentas = pd.read_excel(out_path, sheet_name="Cuentas")
    assert len(df_cuentas) == 1
    assert df_cuentas["Confianza_minima"].iloc[0] == 1.0
