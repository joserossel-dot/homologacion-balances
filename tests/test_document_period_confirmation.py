from io import BytesIO
from types import SimpleNamespace

import pandas as pd
from openpyxl import load_workbook

import app_validacion as app
from extractor_metadata import MetadataEmpresa


def _archivo(nombre: str, contenido: bytes) -> BytesIO:
    archivo = BytesIO(contenido)
    archivo.name = nombre
    return archivo


def _estado_empresa() -> dict:
    return {
        "company_rut": "76.000.000-0",
        "company_razon": "Empresa de prueba",
        "company_giro": "Otro",
        "company_moneda": "$",
        "company_mes": "Diciembre",
        "company_anio": 2026,
        "company_numero_meses": 12,
        "document_pages": {},
        "document_closing_years": {},
        "document_metadata_confirmations": {},
    }


def _confirmar_anio(state: dict, archivo, anio: int | None) -> None:
    state["document_closing_years"][archivo.name] = {
        **app._contexto_confirmacion_anio(archivo),
        "anio_cierre": anio,
        "confirmado": True,
    }


def _confirmar_metadata(
        state: dict, archivo, *, rut: str, razon_social: str, giro: str,
        moneda: str, mes_cierre: str, numero_meses: int,
) -> None:
    state["document_metadata_confirmations"][archivo.name] = {
        **app._contexto_confirmacion_anio(archivo),
        "rut": rut,
        "razon_social": razon_social,
        "giro": giro,
        "moneda": moneda,
        "mes_cierre": mes_cierre,
        "numero_meses": numero_meses,
        "confirmado": True,
    }


def test_detecta_anio_solo_desde_metadata_y_no_completa_ausentes():
    assert app._anio_cierre_detectado(
        MetadataEmpresa(periodo_hasta="31/12/2024"),
    ) == 2024
    assert app._anio_cierre_detectado(MetadataEmpresa()) is None


def test_ui_confirma_anos_por_archivo_con_el_detectado_como_valor_inicial(monkeypatch):
    from streamlit.testing.v1 import AppTest

    monkeypatch.delenv("PILOT_MODE", raising=False)

    def run():
        import streamlit as st
        from io import BytesIO
        import app_validacion as app_local

        primero = BytesIO(b"primero")
        primero.name = "primero.pdf"
        segundo = BytesIO(b"segundo")
        segundo.name = "segundo.pdf"
        previo = app_local._extraer_lineas_encabezado
        app_local._extraer_lineas_encabezado = lambda archivo: [
            "Al 31 de Diciembre de 2024"
            if archivo.name == "primero.pdf"
            else "Al 31 de Diciembre de 2022"
        ]
        try:
            if app_local._confirmar_anios_cierre_documentos([primero, segundo]):
                st.success("años confirmados")
        finally:
            app_local._extraer_lineas_encabezado = previo

    at = AppTest.from_function(run).run()
    assert not at.exception
    assert [select.value for select in at.selectbox] == [2024, 2022]
    at.button[0].click().run()
    assert not at.exception
    assert at.success
    assert at.session_state.document_closing_years["primero.pdf"]["anio_cierre"] == 2024
    assert at.session_state.document_closing_years["segundo.pdf"]["anio_cierre"] == 2022


def test_ui_confirma_metadata_distinta_por_archivo_con_valores_detectados():
    from streamlit.testing.v1 import AppTest

    def run():
        import streamlit as st
        from io import BytesIO
        import app_validacion as app_local

        primero = BytesIO(b"primero")
        primero.name = "primero.pdf"
        segundo = BytesIO(b"segundo")
        segundo.name = "segundo.pdf"
        st.session_state.document_pages = {
            primero.name: [1], segundo.name: [2, 3],
        }
        st.session_state.document_closing_years = {
            primero.name: {
                **app_local._contexto_confirmacion_anio(primero),
                "anio_cierre": 2024, "confirmado": True,
            },
            segundo.name: {
                **app_local._contexto_confirmacion_anio(segundo),
                "anio_cierre": 2022, "confirmado": True,
            },
        }
        previa = app_local._campos_metadata_detectada
        app_local._campos_metadata_detectada = lambda archivo: (
            {
                "rut": "76.000.001-1",
                "razon_social": "Primero SpA",
                "giro": "Inmobiliaria",
                "moneda": "MM",
                "mes_cierre": "Junio",
                "numero_meses": 6,
            }
            if archivo.name == "primero.pdf" else {
                "rut": "77.000.002-2",
                "razon_social": "Segundo Ltda.",
                "giro": "Construcción",
                "moneda": "USD",
                "mes_cierre": "Marzo",
                "numero_meses": 3,
            }
        )
        try:
            if app_local._confirmar_metadata_documentos([primero, segundo]):
                st.success("metadata confirmada")
        finally:
            app_local._campos_metadata_detectada = previa

    at = AppTest.from_function(run).run()
    assert not at.exception
    assert [input_.value for input_ in at.text_input] == [
        "76.000.001-1", "Primero SpA", "Inmobiliaria", "MM",
        "77.000.002-2", "Segundo Ltda.", "Construcción", "USD",
    ]
    assert [select.value for select in at.selectbox] == ["Junio", 6, "Marzo", 3]

    at.button[0].click().run()

    assert not at.exception
    assert at.success
    primero_confirmado = at.session_state.document_metadata_confirmations["primero.pdf"]
    segundo_confirmado = at.session_state.document_metadata_confirmations["segundo.pdf"]
    assert (primero_confirmado["rut"], primero_confirmado["moneda"]) == (
        "76.000.001-1", "MM",
    )
    assert (segundo_confirmado["rut"], segundo_confirmado["moneda"]) == (
        "77.000.002-2", "USD",
    )


def test_aplica_metadata_y_anios_confirmados_por_archivo_sin_cruzarlos(monkeypatch):
    state = _estado_empresa()
    primer_balance = _archivo("balance-2024.pdf", b"primer balance")
    segundo_balance = _archivo("balance-2022.pdf", b"segundo balance")
    state["document_pages"] = {
        primer_balance.name: [1],
        segundo_balance.name: [2, 3],
    }
    monkeypatch.setattr(app.st, "session_state", state)
    _confirmar_anio(state, primer_balance, 2024)
    _confirmar_anio(state, segundo_balance, 2022)
    _confirmar_metadata(
        state, primer_balance,
        rut="76.000.001-1", razon_social="Inmobiliaria Uno SpA",
        giro="Inmobiliaria", moneda="MM", mes_cierre="Junio",
        numero_meses=6,
    )
    _confirmar_metadata(
        state, segundo_balance,
        rut="77.000.002-2", razon_social="Constructora Dos Ltda.",
        giro="Construcción", moneda="USD", mes_cierre="Marzo",
        numero_meses=3,
    )

    meta_2024 = app._aplicar_metadata_confirmada(
        MetadataEmpresa(), archivo=primer_balance,
    )
    meta_2022 = app._aplicar_metadata_confirmada(
        MetadataEmpresa(), archivo=segundo_balance,
    )

    assert meta_2024.anio_cierre == 2024
    assert meta_2024.periodo_desde == "01/01/2024"
    assert meta_2024.periodo_hasta == "30/06/2024"
    assert meta_2022.anio_cierre == 2022
    assert meta_2022.periodo_desde == "01/01/2022"
    assert meta_2022.periodo_hasta == "31/03/2022"
    assert (meta_2024.rut, meta_2024.razon_social, meta_2024.giro) == (
        "76.000.001-1", "Inmobiliaria Uno SpA", "Inmobiliaria",
    )
    assert (meta_2022.rut, meta_2022.razon_social, meta_2022.giro) == (
        "77.000.002-2", "Constructora Dos Ltda.", "Construcción",
    )
    assert meta_2024.moneda == "MM"
    assert meta_2022.moneda == "USD"
    payload_segundo = app._payload_entrega_json(
        archivo_nombre=segundo_balance.name,
        meta=meta_2022,
        estado_entrega={
            "estado": "REVISION", "etiqueta": "Revisión",
            "exportacion_final_permitida": False,
            "exportacion_revision_permitida": True,
            "razones": [],
        },
        emision={"definitivo": False, "motivos": []},
        presentacion=pd.DataFrame(),
        detalle=pd.DataFrame(),
    )
    assert payload_segundo["metadata"] == {
        "razon_social": "Constructora Dos Ltda.",
        "rut": "77.000.002-2",
        "periodo_desde": "01/01/2022",
        "periodo_hasta": "31/03/2022",
        "moneda": "USD",
    }


def test_metadata_heterogenea_no_invalida_lote_si_el_primero_sigue_vigente(monkeypatch):
    state = _estado_empresa()
    primero = _archivo("primero.pdf", b"primero")
    segundo = _archivo("segundo.pdf", b"segundo")
    state["document_pages"] = {primero.name: [1], segundo.name: [2]}
    monkeypatch.setattr(app.st, "session_state", state)
    _confirmar_metadata(
        state, primero,
        rut="76.000.001-1", razon_social="Primero SpA", giro="Inmobiliaria",
        moneda="MM", mes_cierre="Junio", numero_meses=6,
    )
    _confirmar_metadata(
        state, segundo,
        rut="77.000.002-2", razon_social="Segundo Ltda.", giro="Construcción",
        moneda="USD", mes_cierre="Marzo", numero_meses=3,
    )
    state["metadata_confirmada"] = True
    state["metadata_confirmada_contexto"] = app._contexto_metadata_general(primero)
    state["resultados"] = {
        primero.name: pd.DataFrame([{"codigo_clasificado": "AC.01"}]),
        segundo.name: pd.DataFrame([{"codigo_clasificado": "PAT.01"}]),
    }

    assert app._asegurar_metadata_general_vigente([primero, segundo])
    assert set(state["resultados"]) == {primero.name, segundo.name}
    assert state["document_metadata_confirmations"][segundo.name]["moneda"] == "USD"


def test_cambio_de_paginas_invalida_el_anio_y_mantiene_reglas_conservadoras(monkeypatch):
    state = _estado_empresa()
    balance = _archivo("balance.pdf", b"mismo archivo")
    state["document_pages"] = {balance.name: [1]}
    monkeypatch.setattr(app.st, "session_state", state)
    _confirmar_anio(state, balance, 2024)
    _confirmar_metadata(
        state, balance,
        rut="76.000.001-1", razon_social="Empresa Uno", giro="",
        moneda="$", mes_cierre="Diciembre", numero_meses=12,
    )

    assert app._confirmacion_anio_vigente(balance)
    state["document_pages"][balance.name] = [2]

    assert not app._confirmacion_anio_vigente(balance)
    meta = app._aplicar_metadata_confirmada(MetadataEmpresa(), archivo=balance)
    assert meta.anio_cierre is None
    assert meta.periodo_desde is None
    assert meta.periodo_hasta is None
    assert meta.rut == ""
    assert meta.razon_social == ""


def test_cambio_externo_de_paginas_elimina_resultado_y_certificacion(monkeypatch):
    state = _estado_empresa()
    balance = _archivo("balance.pdf", b"mismo archivo")
    state["document_pages"] = {balance.name: [1]}
    monkeypatch.setattr(app.st, "session_state", state)
    state["document_contexts"] = {
        balance.name: app._contexto_confirmacion_anio(balance),
    }
    state["document_scope_confirmed"] = ((
        balance.name,
        app._contexto_confirmacion_anio(balance)["sha256"],
    ),)
    _confirmar_anio(state, balance, 2024)
    _confirmar_metadata(
        state, balance,
        rut="76.000.001-1", razon_social="Empresa Uno", giro="",
        moneda="$", mes_cierre="Diciembre", numero_meses=12,
    )
    state["resultados"] = {
        balance.name: pd.DataFrame([{"codigo_clasificado": "AC.01"}]),
    }
    state["metadata_files"] = {balance.name: MetadataEmpresa(anio_cierre=2024)}
    state["extraction_certifications"] = {balance.name: object()}

    state["document_pages"][balance.name] = [2]
    cambios = app._sincronizar_contextos_documentos([balance])

    assert cambios == {balance.name}
    assert balance.name not in state["resultados"]
    assert balance.name not in state["metadata_files"]
    assert balance.name not in state["extraction_certifications"]
    assert balance.name not in state["document_closing_years"]
    assert balance.name not in state["document_metadata_confirmations"]
    assert state["document_scope_confirmed"] == ((
        balance.name,
        app._contexto_confirmacion_anio(balance)["sha256"],
    ),)
    assert state["document_scope_reconfirmation_required"] is True


def test_invalidar_documento_elimina_solo_su_historial(monkeypatch):
    state = _estado_empresa()
    state["historial_decisiones"] = [
        {"Archivo": "reemplazado.pdf", "Cuenta": "Caja"},
        {"Archivo": "otro.pdf", "Cuenta": "Capital"},
        {"Cuenta": "Entrada sin archivo"},
    ]
    monkeypatch.setattr(app.st, "session_state", state)

    app._invalidar_estado_documento("reemplazado.pdf")

    assert state["historial_decisiones"] == [
        {"Archivo": "otro.pdf", "Cuenta": "Capital"},
        {"Cuenta": "Entrada sin archivo"},
    ]


def test_cambio_no_inicial_requiere_reconfirmar_sin_invalidar_el_otro(monkeypatch):
    state = _estado_empresa()
    primero = _archivo("primero.pdf", b"primero")
    cambiado = _archivo("cambiado.pdf", b"cambiado")
    state["document_pages"] = {
        primero.name: [1], cambiado.name: [1],
    }
    monkeypatch.setattr(app.st, "session_state", state)
    signature = tuple(
        (archivo.name, app._contexto_confirmacion_anio(archivo)["sha256"])
        for archivo in (primero, cambiado)
    )
    state["document_scope_confirmed"] = signature
    state["document_contexts"] = {
        archivo.name: app._contexto_confirmacion_anio(archivo)
        for archivo in (primero, cambiado)
    }
    _confirmar_metadata(
        state, primero,
        rut="76.000.001-1", razon_social="Primero SpA", giro="Inmobiliaria",
        moneda="MM", mes_cierre="Junio", numero_meses=6,
    )
    _confirmar_metadata(
        state, cambiado,
        rut="77.000.002-2", razon_social="Cambiado Ltda.", giro="Construcción",
        moneda="USD", mes_cierre="Marzo", numero_meses=3,
    )
    state["resultados"] = {
        primero.name: pd.DataFrame([{"codigo_clasificado": "AC.01"}]),
        cambiado.name: pd.DataFrame([{"codigo_clasificado": "PAT.01"}]),
    }
    state["extraction_revisions"] = {primero.name: 4, cambiado.name: 2}

    state["document_pages"][cambiado.name] = [2]
    cambios = app._sincronizar_contextos_documentos([primero, cambiado])
    cambios_al_confirmar = app._documentos_con_alcance_modificado(
        signature,
        state["document_scope_confirmed"],
        state["document_pages"],
        {primero.name: [1], cambiado.name: [2]},
    )

    assert cambios == {cambiado.name}
    assert state["document_scope_reconfirmation_required"] is True
    assert state["document_scope_confirmed"] == signature
    assert primero.name in state["resultados"]
    assert state["resultados"][primero.name].iloc[0]["codigo_clasificado"] == "AC.01"
    assert state["extraction_revisions"][primero.name] == 4
    assert state["document_metadata_confirmations"][primero.name]["razon_social"] == "Primero SpA"
    assert cambiado.name not in state["document_metadata_confirmations"]
    assert cambiado.name not in state["resultados"]
    # El segundo envío del formulario conserva el documento que no cambió.
    assert cambios_al_confirmar == set()


def test_cambio_de_origen_reabre_propagaciones_trazadas_y_heredadas(monkeypatch):
    state = _estado_empresa()
    origen = _archivo("origen.pdf", b"origen")
    destino_trazado = "destino-trazado.pdf"
    destino_heredado = "destino-heredado.pdf"
    state["resultados"] = {
        origen.name: pd.DataFrame([{
            "nombre_original": "Caja", "codigo_clasificado": "AC.01",
            "metodo": "validacion_humana", "confianza": 1.0,
            "requiere_revision": False,
        }]),
        destino_trazado: pd.DataFrame([{
            "nombre_original": "Caja", "codigo_clasificado": "AC.01",
            "metodo": "propagado_automático", "confianza": 1.0,
            "requiere_revision": False,
            "propagacion_origen_archivo": origen.name,
        }]),
        destino_heredado: pd.DataFrame([{
            "nombre_original": "Caja", "codigo_clasificado": "AC.01",
            "metodo": "validacion_humana_propagada", "confianza": 1.0,
            "requiere_revision": False,
        }]),
    }
    state["delivery_readiness"] = {
        destino_trazado: {"estado": "LISTO_PARA_INTEGRAR"},
        destino_heredado: {"estado": "LISTO_PARA_INTEGRAR"},
    }
    monkeypatch.setattr(app.st, "session_state", state)

    app._invalidar_estado_documento(origen.name)

    assert origen.name not in state["resultados"]
    for destino in (destino_trazado, destino_heredado):
        fila = state["resultados"][destino].iloc[0]
        assert fila["codigo_clasificado"] == ""
        assert fila["metodo"] == "propagacion_invalidada"
        assert float(fila["confianza"]) == 0.0
        assert bool(fila["requiere_revision"])
        assert destino not in state["delivery_readiness"]


def test_propagacion_manual_registra_su_archivo_de_origen(monkeypatch):
    state = _estado_empresa()
    state["resultados"] = {
        "origen.pdf": pd.DataFrame([{
            "nombre_original": "Caja", "codigo_clasificado": "AC.01",
            "origen_columna": "activo", "monto": 100,
            "confianza": 1.0, "requiere_revision": False,
        }]),
        "destino.pdf": pd.DataFrame([{
            "nombre_original": "Caja", "codigo_clasificado": "",
            "origen_columna": "activo", "monto": 100,
            "confianza": 0.0, "requiere_revision": True,
        }]),
    }
    monkeypatch.setattr(app.st, "session_state", state)

    app.propagar_clasificacion_resultados(
        "Caja", "AC.01", "validacion_humana_propagada",
        archivo_origen="origen.pdf",
    )

    fila = state["resultados"]["destino.pdf"].iloc[0]
    assert fila["codigo_clasificado"] == "AC.01"
    assert fila["propagacion_origen_archivo"] == "origen.pdf"


def test_cambio_del_primer_documento_limpia_residuo_global_y_conserva_otro(monkeypatch):
    state = _estado_empresa()
    fuente_original = _archivo("fuente.pdf", b"version original")
    destino = _archivo("destino.pdf", b"destino")
    state["document_pages"] = {
        fuente_original.name: [1], destino.name: [1],
    }
    monkeypatch.setattr(app.st, "session_state", state)
    state["document_contexts"] = {
        fuente_original.name: app._contexto_confirmacion_anio(fuente_original),
        destino.name: app._contexto_confirmacion_anio(destino),
    }
    state["metadata_confirmada"] = True
    state["metadata_confirmada_contexto"] = app._contexto_metadata_general(
        fuente_original,
    )
    state["resultados"] = {
        fuente_original.name: pd.DataFrame([{"codigo_clasificado": "AC.01"}]),
        destino.name: pd.DataFrame([{"codigo_clasificado": "PAT.01"}]),
    }
    state["extraction_revisions"] = {
        fuente_original.name: 2,
        destino.name: 4,
    }
    fuente_modificada = _archivo("fuente.pdf", b"contenido reemplazado")

    app._sincronizar_contextos_documentos([fuente_modificada, destino])
    vigente = app._asegurar_metadata_general_vigente(
        [fuente_modificada, destino],
    )

    assert not vigente
    assert state["metadata_confirmada"] is False
    assert "metadata_confirmada_contexto" not in state
    for campo in app._CAMPOS_METADATA_GENERAL:
        assert campo not in state
    assert fuente_original.name not in state["resultados"]
    assert destino.name in state["resultados"]
    assert state["resultados"][destino.name].iloc[0]["codigo_clasificado"] == "PAT.01"
    assert state["extraction_revisions"][destino.name] == 4


def test_resultado_historico_usa_el_anio_confirmado_del_documento(monkeypatch):
    state = _estado_empresa()
    documento_historico = _archivo("documento-a.pdf", b"a")
    documento_corriente = _archivo("documento-b.pdf", b"b")
    monkeypatch.setattr(app.st, "session_state", state)
    _confirmar_anio(state, documento_historico, 2024)
    _confirmar_anio(state, documento_corriente, 2022)
    _confirmar_metadata(
        state, documento_historico,
        rut="76.000.001-1", razon_social="Histórica SpA", giro="",
        moneda="$", mes_cierre="Diciembre", numero_meses=12,
    )
    _confirmar_metadata(
        state, documento_corriente,
        rut="77.000.002-2", razon_social="Corriente Ltda.", giro="",
        moneda="$", mes_cierre="Diciembre", numero_meses=12,
    )
    anio_historico = app._aplicar_metadata_confirmada(
        MetadataEmpresa(), archivo=documento_historico,
    ).anio_cierre
    anio_corriente = app._aplicar_metadata_confirmada(
        MetadataEmpresa(), archivo=documento_corriente,
    ).anio_cierre

    contexto = {
        "codigo_original": "2402-12",
        "seccion_contable": "activo",
    }
    assert app._codigo_compatible_con_origen(
        "PAT.03", "activo", 100, "RESULTADO AÑO 2022",
        anio_cierre=anio_historico, **contexto,
    )
    assert not app._codigo_compatible_con_origen(
        "PAT.03", "activo", 100, "RESULTADO AÑO 2022",
        anio_cierre=anio_corriente, **contexto,
    )
    assert app._codigo_compatible_con_origen(
        "PAT.04", "activo", 100, "RESULTADO AÑO 2022",
        anio_cierre=anio_corriente, **contexto,
    )


def test_guardia_de_piloto_bloquea_integracion_y_marca_payload(monkeypatch):
    monkeypatch.setenv("PILOT_MODE", "1")
    estado = {
        "estado": "LISTO_PARA_INTEGRAR",
        "etiqueta": "Listo para integrar",
        "exportacion_final_permitida": True,
        "exportacion_revision_permitida": False,
        "razones": [],
    }

    protegido = app._aplicar_guardia_piloto_entrega(estado)
    payload = app._payload_entrega_json(
        archivo_nombre="balance.pdf",
        meta=None,
        estado_entrega=estado,
        emision={"definitivo": True, "motivos": []},
        presentacion=pd.DataFrame(),
        detalle=pd.DataFrame(),
    )

    assert protegido["estado"] == "LISTO_PARA_INTEGRAR"
    assert protegido["estado_tecnico"] == "LISTO_PARA_INTEGRAR"
    assert protegido["estado_piloto"] == "REVISION_PROVISIONAL"
    assert protegido["etiqueta"] == "Piloto: paquete provisional"
    assert protegido["etiqueta_tecnica"] == "Listo para integrar"
    assert protegido["exportacion_final_permitida"] is False
    assert protegido["exportacion_revision_permitida"] is True
    assert protegido["piloto"] is True
    assert payload["estado_entrega"]["estado"] == "LISTO_PARA_INTEGRAR"
    assert payload["estado_entrega"]["estado_piloto"] == "REVISION_PROVISIONAL"
    assert payload["estado_entrega"]["etiqueta"] == "Piloto: paquete provisional"
    assert payload["emision"]["definitiva"] is False
    assert payload["ejecucion"]["modo_piloto"] is True
    assert payload["ejecucion"]["rama"]
    assert payload["ejecucion"]["commit"]


def test_guardia_de_piloto_no_altera_entrega_fuera_de_piloto(monkeypatch):
    monkeypatch.delenv("PILOT_MODE", raising=False)
    estado = {
        "estado": "LISTO_PARA_INTEGRAR",
        "etiqueta": "Listo para integrar",
        "exportacion_final_permitida": True,
        "exportacion_revision_permitida": False,
        "razones": [],
    }

    assert app._aplicar_guardia_piloto_entrega(estado) is estado


def test_piloto_preserva_diagnostico_tecnico_no_legible(monkeypatch):
    monkeypatch.setenv("PILOT_MODE", "1")
    certificacion = SimpleNamespace(
        estado="fallida",
        metodo="ocr_coordinates_8_amounts",
        orientacion_certificada=False,
        razones=["Sin evidencia de orientación"],
    )

    estado = app._evaluar_estado_entrega_documento(
        pd.DataFrame(), certificacion, {},
    )

    assert estado["estado"] == "NO_LEGIBLE"
    assert estado["estado_tecnico"] == "NO_LEGIBLE"
    assert estado["estado_piloto"] == "SOLO_DIAGNOSTICO"
    assert estado["etiqueta"] == "Piloto: sólo diagnóstico"
    assert estado["exportacion_final_permitida"] is False
    assert estado["exportacion_revision_permitida"] is False


def test_ui_de_piloto_no_ofrece_descarga_definitiva_ni_integracion(monkeypatch):
    from streamlit.testing.v1 import AppTest

    monkeypatch.setenv("PILOT_MODE", "1")
    monkeypatch.setenv("QUALITY_CONTROL_ENFORCE_EXPORT", "false")
    filas = [
        {
            "nombre_original": "Caja", "codigo_original": "1101-01",
            "monto": 100, "origen_columna": "activo",
            "origen_columna_efectiva": "activo", "codigo_clasificado": "AC.01",
            "es_total": False, "requiere_revision": False,
            "metodo": "validacion_humana", "confianza": 1.0,
            "nombre_revision_usuario": "",
        },
        {
            "nombre_original": "Capital", "codigo_original": "2401-01",
            "monto": 100, "origen_columna": "pasivo",
            "origen_columna_efectiva": "patrimonio", "codigo_clasificado": "PAT.01",
            "es_total": False, "requiere_revision": False,
            "metodo": "validacion_humana", "confianza": 1.0,
            "nombre_revision_usuario": "",
        },
    ]

    def run(rows):
        import pandas as pandas_local
        import streamlit as st
        import app_validacion as app_local
        from types import SimpleNamespace as SimpleNamespaceLocal
        from extractor_metadata import MetadataEmpresa
        from reporting_integrity import catalogo_local

        st.session_state.setdefault(
            "metadata_files", {"caso.pdf": MetadataEmpresa(moneda="$")},
        )
        st.session_state.setdefault("extraction_certifications", {
            "caso.pdf": SimpleNamespaceLocal(
                estado="certificada",
                metodo="coordinates_8_amounts",
                orientacion_certificada=True,
                razones=[],
            ),
        })
        previo = app_local.st.download_button

        def capturar(label, **kwargs):
            st.session_state["export_label"] = label
            st.session_state["export_kwargs"] = kwargs

        app_local.st.download_button = capturar
        try:
            app_local._tab_balance(
                pandas_local.DataFrame(rows), catalogo_local(), "caso.pdf",
            )
        finally:
            app_local.st.download_button = previo

    at = AppTest.from_function(run, args=(filas,)).run(timeout=20)
    assert not at.exception
    assert "piloto provisional" in at.session_state.export_label.lower()
    descarga = at.session_state.export_kwargs
    assert descarga["file_name"].startswith("PILOTO-")
    libro = load_workbook(BytesIO(descarga["data"]), data_only=True)
    assert libro["Balance Normalizado"]["F1"].value == (
        "PILOTO: PAQUETE PROVISIONAL, NO INTEGRAR"
    )
    controles = {
        fila[0]: fila[1]
        for fila in libro["Control de emisión"].iter_rows(min_row=2, values_only=True)
        if fila[0]
    }
    assert controles["Modo piloto"] is True
    assert controles["Rama de ejecución"]
    assert controles["Commit de ejecución"]
